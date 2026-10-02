from __future__ import annotations

from copy import copy

from openpyxl.formula.translate import Translator
from openpyxl.worksheet.worksheet import Worksheet

from ..config import cfg


def sort_bom_sections(ws: Worksheet) -> None:
    """只在各連續料號區塊內排序，保留標題、空白分隔及整列格式。"""
    part_col = cfg("excel.bom_part_col", 2) + 1
    data_start = cfg("excel.bom_data_start_row", 5)
    merged_ranges = list(ws.merged_cells.ranges)
    vertical_ranges = [area for area in merged_ranges if area.min_row != area.max_row]
    boundaries = {row for area in vertical_ranges for row in (area.min_row, area.max_row + 1)}
    headings = {"料號", "自備", "自備料", "客供", "客供料"}
    sections: list[list[int]] = []
    current: list[int] = []
    for row in range(data_start, ws.max_row + 1):
        part = str(ws.cell(row, part_col).value or "").strip()
        merged_part = any(
            area.min_row <= row <= area.max_row and area.min_col <= part_col <= area.max_col
            for area in merged_ranges
        )
        if row in boundaries or not part or part in headings or merged_part:
            if current:
                sections.append(current)
                current = []
        if part and part not in headings and not merged_part:
            current.append(row)
    if current:
        sections.append(current)

    for rows in sections:
        ordered = sorted(rows, key=lambda row: str(ws.cell(row, part_col).value).strip().upper())
        if rows == ordered:
            continue
        destinations = dict(zip(ordered, rows))
        cells = {row: [copy(cell) for cell in ws[row]] for row in rows}
        dimensions = {row: copy(ws.row_dimensions[row]) for row in rows}
        row_merges = [area for area in merged_ranges if area.min_row == area.max_row and area.min_row in destinations]
        for area in row_merges:
            ws.unmerge_cells(str(area))

        for source_row, target_row in destinations.items():
            for source_cell in cells[source_row]:
                column = source_cell.column
                # 垂直合併的區塊標籤固定原位，只有區塊內的料件明細移動。
                if any(area.min_row <= target_row <= area.max_row and area.min_col <= column <= area.max_col for area in vertical_ranges):
                    continue
                target_cell = ws.cell(target_row, column)
                value = source_cell.value
                if source_cell.data_type == "f" and isinstance(value, str):
                    value = Translator(value, origin=source_cell.coordinate).translate_formula(target_cell.coordinate)
                target_cell.value = value
                target_cell.data_type = source_cell.data_type
                target_cell._style = copy(source_cell._style)
                target_cell.comment = copy(getattr(source_cell, "comment", None))
                target_cell._hyperlink = copy(getattr(source_cell, "hyperlink", None))
                if target_cell.hyperlink:
                    target_cell.hyperlink.ref = target_cell.coordinate
            dimension = dimensions[source_row]
            dimension.index = target_row
            ws.row_dimensions[target_row] = dimension

        for area in row_merges:
            target_row = destinations[area.min_row]
            ws.merge_cells(start_row=target_row, end_row=target_row, start_column=area.min_col, end_column=area.max_col)
