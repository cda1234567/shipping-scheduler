"""在歷史主檔批次之前插欄，保留後方事件並重算結存。"""
from __future__ import annotations

from copy import copy
import re

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

from .main_file_recalc import _stock_events, recalc_batch_balances_for_cell
from ..models import calc_suggested_qty

_CODE = re.compile(r"^(\d+)-(\d+)$")
_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?\d*)$")


def _shift_reference(value: str, start: int, width: int, sheet: str, formula_sheet: str) -> str:
    prefix, sep, address = value.rpartition("!")
    if sep:
        if prefix.strip("'").replace("''", "'") != sheet:
            return value
    else:
        address = value
        if formula_sheet != sheet:
            return value
    parts = []
    for endpoint in address.split(":"):
        match = _CELL.fullmatch(endpoint)
        if not match:
            return value
        absolute, letters, row = match.groups()
        col = column_index_from_string(letters)
        parts.append(f"{absolute}{get_column_letter(col + width if col >= start else col)}{row}")
    return (prefix + "!" if sep else "") + ":".join(parts)


def insert_columns(ws, start: int, width: int) -> None:
    """openpyxl insert_cols 不更新公式、欄寬或合併範圍，逐項同步。"""
    formulas = []
    for sheet in ws.parent.worksheets:
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                tokens = Tokenizer(cell.value).items
                for token in tokens:
                    if token.type == "OPERAND" and token.subtype == "RANGE":
                        token.value = _shift_reference(token.value, start, width, ws.title, sheet.title)
                formulas.append((sheet, cell.row, cell.column, "=" + "".join(t.value for t in tokens)))
    merged = [copy(item) for item in ws.merged_cells.ranges]
    for item in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(item))
    dimensions = [(key, copy(value)) for key, value in ws.column_dimensions.items()]
    ws.insert_cols(start, width)
    ws.column_dimensions.clear()
    for key, dimension in dimensions:
        old_col = column_index_from_string(key)
        new_col = old_col + width if old_col >= start else old_col
        dimension.index = get_column_letter(new_col)
        if dimension.min and dimension.min >= start:
            dimension.min += width
        if dimension.max and dimension.max >= start:
            dimension.max += width
        ws.column_dimensions[dimension.index] = dimension
    for item in merged:
        if item.min_col >= start:
            item.shift(col_shift=width)
        elif item.max_col >= start:
            item.max_col += width
        ws.merge_cells(str(item))
    for sheet, row, col, formula in formulas:
        if sheet is ws and col >= start:
            col += width
        sheet.cell(row, col).value = formula


def validate_insertion_anchor(ws, anchor: str) -> int:
    match = _CODE.fullmatch(anchor)
    if not match:
        raise ValueError("插入位置必須是完整單據編號，例如 10-0")
    starts = [col for col in range(1, ws.max_column + 1)
              if str(ws.cell(1, col).value or "").strip() == anchor]
    if not starts:
        raise ValueError(f"主檔找不到插入位置 {anchor}，請重新選擇")
    start = min(starts)
    # 同單多份 BOM 允許相鄰的三欄組；分散重複編號無法安全定位。
    if starts != list(range(start, start + 3 * len(starts), 3)):
        raise ValueError(f"主檔編號 {anchor} 出現在不連續位置，無法插入")
    return start


def build_insertion_plan(ws, batches, decisions, moq_map, st_inventory_stock, build_plan, write_plan):
    result = {"batches": [], "shortages": [], "merged_parts": 0, "code_shifts": {}}
    for batch in batches:
        anchor = str(batch.get("insert_before_code") or "").strip()
        if not anchor:
            plan = build_plan(ws, [batch], decisions, moq_map, st_inventory_stock)
            write_plan(ws.parent, plan)
        else:
            start = validate_insertion_anchor(ws, anchor)
            match = _CODE.fullmatch(anchor)
            prefix_ws = ws.parent.copy_worksheet(ws)
            try:
                prefix_ws.delete_cols(start, prefix_ws.max_column - start + 1)
                plan = build_plan(prefix_ws, [batch], decisions, moq_map, st_inventory_stock)
                # 新料列由規劃器建立，保留其資料與樣式。
                for row in range(2, prefix_ws.max_row + 1):
                    if prefix_ws.cell(row, 1).value == ws.cell(row, 1).value:
                        continue
                    ws.insert_rows(row)
                    for prior in result["batches"]:
                        for group in prior.get("groups") or []:
                            for item in group.get("rows") or []:
                                if item["row_idx"] >= row:
                                    item["row_idx"] += 1
                    for col in range(1, start):
                        source = prefix_ws.cell(row, col)
                        target = ws.cell(row, col)
                        target.value, target._style = source.value, copy(source._style)
            finally:
                ws.parent.remove(prefix_ws)
            width = 3 * sum(len(item.get("groups") or []) for item in plan["batches"])
            if not width:
                raise ValueError("這筆單據沒有可寫入主檔的用料")
            shifts = {}
            for col in range(start, ws.max_column + 1):
                text = str(ws.cell(1, col).value or "").strip()
                code = _CODE.fullmatch(text)
                if code and code[1] == match[1] and int(code[2]) >= int(match[2]):
                    shifts[text] = f"{code[1]}-{int(code[2]) + 1}"
            for col in range(start, ws.max_column + 1):
                cell = ws.cell(1, col)
                text = str(cell.value or "").strip()
                if text in shifts:
                    cell.value = shifts[text]
            insert_columns(ws, start, width)
            # 較早處理的同批新單也可能被後來插入的單據推移。
            for prior in result["batches"]:
                for group in prior.get("groups") or []:
                    for value in [group] + list(group.get("rows") or []):
                        for key in ("col_h", "col_f", "col_j"):
                            if value.get(key, 0) >= start:
                                value[key] += width
            write_plan(ws.parent, plan)
            affected_rows = {item["row_idx"] for group in plan["batches"][0]["groups"] for item in group["rows"]}
            events = _stock_events(ws)
            known_columns = {col for event in events
                             for col in range(int(event["start_col"]), int(event["balance_col"]) + 1)}
            for row in affected_rows:
                for col in range(start + width, ws.max_column + 1):
                    cell = ws.cell(row, col)
                    if col not in known_columns and (isinstance(cell.value, (int, float)) or cell.data_type == "f"):
                        raise ValueError("後續存在無法辨識的庫存欄位，請先補齊批次編號後再插入")
                # 需求/補料公式無法由 openpyxl 求值，禁止把公式當零寫壞歷史結存。
                for event in events:
                    if int(event["start_col"]) < start + width:
                        continue
                    for col in range(int(event["start_col"]), int(event["balance_col"])):
                        if ws.cell(row, col).data_type == "f":
                            raise ValueError("後續補料或用量含公式，請先將該數量確認為數值後再插入")
                recalc_batch_balances_for_cell(ws, row=row, col=start)
                tail_values = [float(ws.cell(row, int(event["balance_col"])).value)
                               for event in events if int(event["start_col"]) >= start + width
                               and isinstance(ws.cell(row, int(event["balance_col"])).value, (int, float))]
                if tail_values and min(tail_values) < 0:
                    part = str(ws.cell(row, 1).value or "").strip().upper()
                    shortage = -min(tail_values)
                    moq = float((moq_map or {}).get(part) or 0)
                    plan["shortages"].append({
                        "order_id": batch.get("order_id"), "part_number": part,
                        "batch_code": anchor, "model": batch.get("model", ""),
                        "description": str(ws.cell(row, 4).value or ""),
                        "shortage_amount": shortage, "resulting_stock": -shortage,
                        "current_stock": 0, "needed": shortage, "moq": moq,
                        "supplement_qty": float((batch.get("supplements") or {}).get(part) or 0),
                        "suggested_qty": calc_suggested_qty(shortage, moq),
                        "decision": (batch.get("decisions") or decisions or {}).get(part, "None"),
                        "downstream": True,
                    })
            plan["batches"][0]["code_shifts"] = shifts
            result["code_shifts"].update(shifts)
        result["batches"].extend(plan["batches"])
        result["shortages"].extend(plan["shortages"])
        result["merged_parts"] += plan["merged_parts"]
    # 插入後的核對應使用最末結存，不能拿中途新單的結存當目前庫存。
    from .merge_to_main import _read_latest_stock
    for batch_plan in result["batches"]:
        for group in batch_plan.get("groups") or []:
            for row in group.get("rows") or []:
                row["j_value"] = ws.cell(row["row_idx"], row["col_j"]).value
                row["current_stock"] = row["j_value"] - row["effective_h"] + row["f_value"]
                # 插入重算以已寫入的整數補料與用量為準。
                row["effective_h_raw"] = row["effective_h"]
                row["needed_qty_raw"] = row["f_value"]
                row["final_stock"] = _read_latest_stock(ws, row["row_idx"], ws.max_column)
    result["new_col_count"] = ws.max_column
    result["already_written"] = True
    return result
