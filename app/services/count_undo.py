"""主檔盤點的受限撤回：只接受伺服器保存的備份證據。"""
from __future__ import annotations

import hashlib
import hmac
import json
from datetime import datetime
from pathlib import Path
import re

from .. import database as db
from ..config import BACKUP_DIR, DATA_DIR
from .inventory_restore_guard import _read_count_boundaries
from .main_file_lock import serialized_main_file_write
from .main_reader import read_stock, read_moq
from .merge_to_main import backup_main_file


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _digest(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _state_hash(conn) -> str:
    return _digest(db.main_count_undo_state(conn))


def make_count_undo_receipt(main_path: str, backup_path: str) -> dict:
    with db.get_conn() as conn:
        fingerprint = _state_hash(conn)
    return dict(backup_name=Path(backup_path).name, backup_sha256=_sha(Path(backup_path)),
                main_sha256=_sha(Path(main_path)), main_path=str(Path(main_path).resolve()),
                period_id=db.get_current_main_period_id(), db_fingerprint=fingerprint, restored_at='')


def _backup_path(name: str) -> Path:
    root = Path(BACKUP_DIR).resolve()
    path = (root / name).resolve()
    if not name or Path(name).name != name or '\\' in name or path.parent != root or path.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise ValueError('盤點備份路徑不合法')
    if not path.is_file():
        raise ValueError('找不到這次盤點前的主檔備份')
    return path


def _backup_time(name: str) -> str:
    match = re.search(r'_backup_(\d{8}_\d{6})', name)
    return datetime.strptime(match[1], '%Y%m%d_%H%M%S').isoformat() if match else ''


def _legacy_receipt(session: dict, main: Path, conn) -> dict:
    # 舊版清理報告已記錄前後雜湊；不從名稱或最近檔案猜測還原目標。
    report_path = Path(DATA_DIR) / f"count{int(session['id'])}-repair-preview-new.json"
    if not report_path.is_file():
        raise ValueError('舊盤點缺少安全還原憑證，需先核對備份')
    report = json.loads(report_path.read_text(encoding='utf-8'))
    if report.get('session_id') != session['id'] or any(report.get(k) != 0 for k in ('stock_changed_count', 'unselected_cells_changed', 'dispatch_supply_usage_changed')):
        raise ValueError('舊盤點修復憑證不完整，無法安全撤回')
    completed = datetime.fromisoformat(session['completed_at'])
    matches = []
    for path in Path(BACKUP_DIR).glob('*_backup_*.xls*'):
        stamp = _backup_time(path.name)
        if stamp and 0 <= (completed - datetime.fromisoformat(stamp)).total_seconds() <= 120 and _sha(path) == report.get('reference_sha256'):
            matches.append(path)
    if len(matches) != 1:
        raise ValueError('找不到可唯一核對的盤點前備份')
    reason = ''
    # 原始工作階段仍保留；主檔標記及完整雜湊證明報告確實屬於本次盤點。
    stat = main.stat()
    boundaries = _read_count_boundaries(str(main), stat.st_mtime_ns, stat.st_size)
    matching = [b for b in boundaries if b.get('session_id') == session['id'] and b.get('cutoff_at') == session['cutoff_at']]
    if len(matching) != 1 or sorted(matching[0].get('absorbed_record_ids', [])) != sorted(report.get('record_ids', [])):
        reason = '目前主檔已不是這次盤點結果'
    boundary = matching[0] if matching else {}
    if boundary and conn.execute('SELECT COALESCE(MAX(id),0) FROM defective_records').fetchone()[0] > boundary['max_record_id']:
        reason = '盤點後已有新增不良品／多打，不能直接撤回'
    if boundary:
        # 舊版沒有 DB 指紋，重新核對當時清理的精確紀錄與數量；不能把現在的異動當成原始基準。
        import openpyxl
        from .count_absorption import plan_count_absorption
        from .main_insertion import validate_insertion_anchor
        workbook = openpyxl.load_workbook(matches[0], keep_vba=matches[0].suffix.lower() == '.xlsm')
        try:
            ws = workbook.worksheets[0]
            anchor = validate_insertion_anchor(ws, session['cutoff_code'])
            groups = sum(str(ws.cell(1, c).value or '').strip() == session['cutoff_code'] for c in range(1, ws.max_column + 1))
            records = [dict(r) for r in conn.execute('SELECT r.*,b.main_period_id FROM defective_records r LEFT JOIN defective_batches b ON b.id=r.batch_id')]
            plan = plan_count_absorption(ws, anchor + 3 * groups, boundary['parts'], records,
                                         cutoff_at=boundary['cutoff_at'], max_record_id=boundary['max_record_id'], current_period_id=_period(conn))
            if plan['record_ids'] != sorted(report.get('record_ids', [])):
                reason = '已併入的不良品／多打紀錄與原備份不符，不能安全撤回'
        except ValueError:
            reason = '已併入的不良品／多打數量與原備份不符，不能安全撤回'
        finally:
            workbook.close()
    if conn.execute("SELECT 1 FROM activity_logs WHERE created_at>? AND action NOT IN ('inventory_count_completed','database_backup_created') LIMIT 1", (session['completed_at'],)).fetchone():
        reason = '盤點後已有其他操作，舊盤點無法安全撤回'
    return dict(backup_name=matches[0].name, backup_sha256=report['reference_sha256'],
                main_sha256=report['output_sha256'], main_path=str(main.resolve()),
                period_id=_period(conn), db_fingerprint=_state_hash(conn), restored_at='', legacy_reason=reason)


def _setting(conn, key: str, default=''):
    row = conn.execute('SELECT value FROM settings WHERE key=?', (key,)).fetchone()
    return row[0] if row else default


def _period(conn) -> int:
    return int(_setting(conn, 'main_file_period_id', '0') or 0)


def _inspect(session_id: int, conn) -> tuple[dict, dict, Path | None]:
    session_row = conn.execute("SELECT * FROM inventory_count_sessions WHERE id=? AND scope='st' AND status='completed'", (session_id,)).fetchone()
    if not session_row:
        raise ValueError('找不到已完成的主檔盤點')
    session = dict(session_row)
    item = dict(session_id=session_id, completed_at=session['completed_at'], cutoff_code=session['cutoff_code'],
                backup_name='', backup_at='', can_undo=False, can_download=False, reason='', token='', restored_at='')
    receipt, backup = {}, None
    try:
        main = Path(_setting(conn, 'main_file_path') or '')
        if not main.is_file():
            raise ValueError('找不到目前主檔')
        row = conn.execute('SELECT value FROM settings WHERE key=?', (f'main_count_undo_{session_id}',)).fetchone()
        receipt = json.loads(row[0]) if row else _legacy_receipt(session, main, conn)
        item.update(backup_name=receipt['backup_name'], backup_at=_backup_time(receipt['backup_name']), restored_at=receipt.get('restored_at', ''))
        backup = _backup_path(receipt['backup_name'])
        if _sha(backup) != receipt['backup_sha256']:
            raise ValueError('盤點前備份內容已變更，不能安全還原')
        item['can_download'] = True
        if receipt.get('legacy_reason'):
            raise ValueError(receipt['legacy_reason'])
        if receipt.get('restored_at'):
            raise ValueError('這次盤點已撤回，不能重複撤回')
        if conn.execute("SELECT 1 FROM inventory_count_sessions WHERE status='active' AND scope='st'").fetchone():
            raise ValueError('有盤點正在進行，請先完成或取消')
        if conn.execute("SELECT 1 FROM inventory_count_sessions WHERE status='completed' AND scope='st' AND id>?", (session_id,)).fetchone():
            raise ValueError('已有較新的盤點，不能撤回較舊盤點')
        if str(main.resolve()) != receipt['main_path'] or _period(conn) != receipt['period_id']:
            raise ValueError('主檔或年度已切換，不能撤回舊盤點')
        if _sha(main) != receipt['main_sha256']:
            raise ValueError('盤點後主檔已有發料或其他修改，不能直接覆蓋；請先下載備份核對')
        fingerprint = _state_hash(conn)
        if fingerprint != receipt['db_fingerprint']:
            raise ValueError('盤點後不良品／發料紀錄已有異動，不能直接撤回')
        item['token'] = _digest([session_id, receipt, fingerprint])
        item['can_undo'] = True
    except (ValueError, OSError, KeyError, TypeError) as error:
        item['reason'] = str(error)
    return item, receipt, backup


def list_count_undo() -> dict:
    with db.get_conn() as conn:
        ids = [row[0] for row in conn.execute("SELECT id FROM inventory_count_sessions WHERE scope='st' AND status='completed' ORDER BY id DESC LIMIT 20")]
        return {'items': [_inspect(session_id, conn)[0] for session_id in ids]}


def count_undo_backup(session_id: int) -> Path:
    with db.get_conn() as conn:
        item, _, backup = _inspect(session_id, conn)
    if not item['can_download'] or backup is None:
        raise ValueError(item['reason'] or '備份無法下載')
    return backup


@serialized_main_file_write
def undo_main_count(session_id: int, token: str) -> dict:
    from .st_reconcile import _restore_main_from_backup
    from ..routers.main_file import invalidate_main_data_cache
    main, safety, replaced = None, None, False
    try:
        with db.get_conn() as conn:
            conn.execute('BEGIN IMMEDIATE')
            item, receipt, backup = _inspect(session_id, conn)
            if not item['can_undo']:
                raise ValueError(item['reason'])
            if not token or not hmac.compare_digest(token, item['token']):
                raise ValueError('盤點或主檔狀態已變更，請重新整理後再撤回')
            main = Path(receipt['main_path'])
            stock, moq = read_stock(str(backup)), read_moq(str(backup))
            if not stock:
                raise ValueError('盤點備份無法讀取有效庫存，未寫入任何資料')
            safety = Path(backup_main_file(str(main), str(BACKUP_DIR)))
            if _sha(main) != receipt['main_sha256'] or _sha(backup) != receipt['backup_sha256']:
                raise ValueError('備份期間主檔已變更，未執行撤回')
            _restore_main_from_backup(str(backup), str(main))
            replaced = True
            if _sha(main) != receipt['backup_sha256']:
                raise RuntimeError('還原主檔驗證失敗')
            receipt['safety_backup_name'] = safety.name
            db.apply_main_count_undo(conn, session_id, receipt, stock, moq)
    except Exception:
        if replaced:
            _restore_main_from_backup(str(safety), str(main))
        invalidate_main_data_cache()
        raise
    invalidate_main_data_cache()
    return dict(ok=True, session_id=session_id, restored_at=receipt['restored_at'],
                safety_backup_name=safety.name, snapshot_count=len(stock))
