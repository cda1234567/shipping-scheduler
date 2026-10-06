"""加工廠盤點對帳試算。"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date
import hashlib
import hmac
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from app import database as db
from app.config import BACKUP_DIR
from app.constants import ST_RECONCILE_ADJUSTMENT_REASON
from app.snapshot_sync import refresh_snapshot_from_main

from .reconcile_core import theoretical_stock_with_details
from .main_reader import read_batch_stock, read_stock
from .main_file_lock import serialized_main_file_write
from .main_file_recalc import _stock_events, recalc_batch_balances_for_cell
from .main_insertion import insert_columns, validate_insertion_anchor
from .merge_to_main import PART_COL, _save_workbook_atomically, backup_main_file
from .overrun_deduction import suggest_main_part_numbers
from .xls_reader import open_workbook_any

CUSTOMER_HEADER = "客戶編號"
PART_HEADER = "汎翊國際料號"
DESC_HEADER = "品名規格"
PHYSICAL_HEADER = "辰尚填寫"
GENLIN_PART_HEADER = "consign invoice NO"
GENLIN_BOOK_HEADER = "辰尚庫存"
GENLIN_PHYSICAL_HEADER = "庚霖庫存 當下實際"
GENLIN_DESC_HEADER = "Parts No/Description"
GENLIN_TOTAL_PHYSICAL_HEADER = "實際庫存總和"
GENLIN_PRODUCTION_BOOK_HEADER = "生產結餘"

CATEGORY_HAVE_OURS_NOT_THEIRS = "我有單他沒有"
CATEGORY_HAVE_THEIRS_NOT_OURS = "他有單我沒入"
CATEGORY_QTY_MISMATCH = "同單數量不符"
CATEGORY_UNATTRIBUTED = "無法歸因淨差"
CATEGORY_MATCHED = "無差異"
CATEGORY_STOP_LOSS = "停損吸收"
CATEGORY_GENLIN_BLANK_PHYSICAL = "未填實盤，跳過"


class StaleReconcilePreviewError(ValueError):
    """盤點試算所依據的檔案或資料已變更。"""


class PartMappingError(ValueError):
    """盤點來源料號與主檔料號的人工對應不合法。"""

ASSUMPTIONS = [
    "本試算只讀取上傳盤點表，不會寫入 ST 庫存，也不會建立對齊點。",
    "盤點表沒有良品 / 不良品分欄，本版先用盤點數與系統理論良品庫存比對；差額需人工再對照未報廢不良品單。",
    "盤點表沒有工單或 MO 號，本版只能做料號級淨差歸因；同單數量不符需等盤點表提供單號後才能精準判定。",
    "H 欄視為客戶群組總盤點數；同一群組若有多個汎翊料號，會標示需人工拆分並歸入無法歸因淨差。",
]
def _normalize_text(value: Any) -> str:
    return str(value or "").strip()


def _normalize_header_text(value: Any) -> str:
    return " ".join(_normalize_text(value).split())


def _compact_header_text(value: Any) -> str:
    return "".join(_normalize_text(value).split())


def _normalize_part(value: Any) -> str:
    return _normalize_text(value).upper()


def _normalize_genlin_part(value: Any) -> str:
    return _normalize_part(value)


def _sha256_file(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _display_number(value: float) -> int | float:
    return int(value) if value == int(value) else value


def _resolve_part(part: str, available: set[str]) -> str:
    if part in available:
        return part
    alias = part[:-4] if part.endswith('-TAB') else part + '-TAB'
    return alias if alias in available else part


def _detect_count_date(ws, header_row: int, filename: str) -> str:
    titles = [str(cell or '') for row in ws.iter_rows(min_row=1, max_row=max(1, header_row - 1), values_only=True) for cell in row]
    for text in titles + [filename]:
        for match in re.finditer(r'(?<!\d)(20\d{2})[/-](\d{1,2})[/-](\d{1,2})(?!\d)', text):
            try:
                return date(*(int(value) for value in match.groups())).isoformat()
            except ValueError:
                continue
    return ''


def _normalize_part_numbers(values: list[str] | None) -> list[str] | None:
    if values is None:
        return None
    normalized = [
        part
        for value in values
        for part in [_normalize_part(value)]
        if part and part != "[]"
    ]
    return list(dict.fromkeys(normalized))


def normalize_reconcile_part_mappings(values: dict[str, str] | None) -> dict[str, str]:
    if values is None:
        return {}
    if not isinstance(values, dict):
        raise PartMappingError("料號對應格式需為原料號與主檔料號的對照表")
    normalized: dict[str, str] = {}
    for source, target in values.items():
        if not isinstance(source, str) or not isinstance(target, str):
            raise PartMappingError("料號對應的原料號與主檔料號都需為文字")
        source_part, target_part = _normalize_part(source), _normalize_part(target)
        if not source_part or not target_part:
            raise PartMappingError("料號對應的原料號與主檔料號不可空白")
        if source_part in normalized:
            raise PartMappingError(f"原料號 {source_part} 有重複的料號對應")
        normalized[source_part] = target_part
    return normalized


def _try_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _try_optional_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _find_chenshang_header_row(ws) -> tuple[int, dict[str, int]]:
    required = {CUSTOMER_HEADER, PART_HEADER, PHYSICAL_HEADER}
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True), start=1):
        values = [_normalize_text(cell) for cell in row]
        if not required.issubset(set(values)):
            continue
        return row_idx, {
            "customer": values.index(CUSTOMER_HEADER),
            "part": values.index(PART_HEADER),
            "desc": values.index(DESC_HEADER) if DESC_HEADER in values else -1,
            "physical": values.index(PHYSICAL_HEADER),
        }
    raise ValueError("找不到加工廠盤點表頭，需包含「客戶編號」、「汎翊國際料號」與「辰尚填寫」")


def _find_genlin_header_row(ws) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True), start=1):
        values = [_normalize_header_text(cell) for cell in row]
        compact_values = [_compact_header_text(cell) for cell in row]
        part_col = 3 if len(values) > 3 and values[3].lower().startswith("consign") else -1
        total_physical_col = next(
            (idx for idx, value in enumerate(compact_values) if value.startswith(GENLIN_TOTAL_PHYSICAL_HEADER)),
            -1,
        )
        production_book_col = next(
            (idx for idx, value in enumerate(compact_values) if value.startswith(GENLIN_PRODUCTION_BOOK_HEADER)),
            -1,
        )
        book_col = values.index(GENLIN_BOOK_HEADER) if GENLIN_BOOK_HEADER in values else -1
        physical_col = next(
            (idx for idx, value in enumerate(values) if value.startswith(GENLIN_PHYSICAL_HEADER.split()[0])),
            -1,
        )
        has_detail = book_col >= 0 and physical_col >= 0 and any(
            _cell(data_row, part_col) and any(_cell(data_row, col) not in (None, '') for col in (book_col, physical_col))
            for data_row in ws.iter_rows(min_row=row_idx + 2, values_only=True)
        )
        if not has_detail and total_physical_col >= 0 and production_book_col >= 0:
            book_col, physical_col = production_book_col, total_physical_col
        if part_col < 0 or book_col < 0 or physical_col < 0:
            continue
        # 庚霖實盤固定讀 G 欄，空白時也不可改用總庫存。
        physical_col = 6
        desc_col = next(
            (
                idx
                for idx, value in enumerate(compact_values)
                if value.lower()
                in {_compact_header_text(GENLIN_DESC_HEADER).lower(), "partsno/partsdescription"}
            ),
            -1,
        )
        return row_idx, {
            "part": part_col,
            "desc": desc_col,
            "book": book_col,
            "physical": physical_col,
        }
    raise ValueError("找不到庚霖實際庫存表頭，需包含 D 欄「consign」開頭、「辰尚庫存」與「庚霖庫存」開頭")


def _detect_format(ws) -> tuple[str, int, dict[str, int]]:
    header_errors: list[str] = []
    try:
        header_row, columns = _find_genlin_header_row(ws)
        return "genlin", header_row, columns
    except ValueError as error:
        header_errors.append(str(error))
    try:
        header_row, columns = _find_chenshang_header_row(ws)
        return "chenshang", header_row, columns
    except ValueError as error:
        header_errors.append(str(error))
    raise ValueError("無法辨識盤點表格式；" + "；".join(header_errors))


def _cell(row: tuple[Any, ...], column: int) -> Any:
    if column < 0 or len(row) <= column:
        return None
    return row[column]


def _finish_group(group: dict[str, Any] | None, rows: list[dict]) -> None:
    if not group:
        return
    unique_parts = list(dict.fromkeys(group["parts"]))
    if not unique_parts:
        return

    descriptions = group["descriptions"]
    if len(unique_parts) == 1:
        part = unique_parts[0]
        rows.append({
            "part_number": part,
            "description": descriptions.get(part, ""),
            "physical": float(group["physical"]),
            "customer_code": group["customer_code"],
            "group_part_count": 1,
            "group_parts": unique_parts,
            "needs_manual_split": False,
        })
        return

    for part in unique_parts:
        rows.append({
            "part_number": part,
            "description": descriptions.get(part, ""),
            "physical": None,
            "group_physical": float(group["physical"]),
            "customer_code": group["customer_code"],
            "group_part_count": len(unique_parts),
            "group_parts": unique_parts,
            "needs_manual_split": True,
        })


def _parse_genlin_sheet(ws, header_row: int, columns: dict[str, int]) -> list[dict]:
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
        part = _normalize_genlin_part(_cell(row, columns["part"]))
        if not part:
            continue
        rows.append({
            "part_number": part,
            "description": _normalize_text(_cell(row, columns["desc"])),
            "book_qty": _try_float(_cell(row, columns["book"])),
            "physical_qty": _try_optional_float(_cell(row, columns["physical"])),
        })
    return rows


def _genlin_source_columns(columns: dict[str, int]) -> dict[str, str]:
    return {
        "book": get_column_letter(columns["book"] + 1),
        "physical": get_column_letter(columns["physical"] + 1),
    }


def _build_genlin_assumptions(source_columns: dict[str, str]) -> list[str]:
    book_column = source_columns.get("book") or "F"
    physical_column = source_columns.get("physical") or "G"
    return [
        f"本試算讀取庚霖實際庫存格式：{book_column} 欄為我方帳面，{physical_column} 欄為庚霖實盤。",
        f"停損點模式只分「無差異」與「停損吸收」；{book_column} 與 {physical_column} 的差額不再逐筆歸因。",
        f"按下設為停損點後，系統會以庚霖實盤 {physical_column} 欄重設每個料號的 ST 庫存基準，差額自動寫入調帳紀錄。",
        "報告最後的「未被盤點覆蓋」清單＝這段期間有不良品/多打扣帳、但不在盤點檔裡的料號（多為自備料）——這些料的帳沒有被實盤驗證，僅供你判斷是否請加工廠一併盤點。",
    ]


def parse_st_reconcile_file(path: str) -> dict[str, Any]:
    """解析盤點表；依表頭自動偵測辰尚舊格式或庚霖實際庫存格式。"""
    source_path = Path(path)
    workbook = open_workbook_any(str(source_path), read_only=True, data_only=True)
    try:
        sheet_by_stripped_name = {str(name).strip(): name for name in workbook.sheetnames}
        sheet_name = sheet_by_stripped_name.get("實際庫存-生產結餘")
        ws = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        file_format, header_row, columns = _detect_format(ws)
        if file_format == "genlin":
            parsed_rows = _parse_genlin_sheet(ws, header_row, columns)
            return {
                "format": "genlin",
                "sheet_name": ws.title.strip(),
                "source_columns": _genlin_source_columns(columns),
                "count_date": _detect_count_date(ws, header_row, source_path.name),
                "source_sha256": _sha256_file(str(source_path)),
                "rows": parsed_rows,
                "part_count": len(parsed_rows),
                "manual_split_count": 0,
            }

        current_group: dict[str, Any] | None = None
        parsed_rows: list[dict] = []

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            customer = _normalize_text(_cell(row, columns["customer"]))
            raw_physical = _cell(row, columns["physical"])
            starts_group = bool(customer or raw_physical not in (None, ""))
            if starts_group:
                _finish_group(current_group, parsed_rows)
                current_group = {
                    "customer_code": customer,
                    "physical": _try_float(raw_physical),
                    "parts": [],
                    "descriptions": {},
                }

            part = _normalize_part(_cell(row, columns["part"]))
            if not part:
                continue
            if current_group is None:
                current_group = {
                    "customer_code": "",
                    "physical": 0.0,
                    "parts": [],
                    "descriptions": {},
                }
            current_group["parts"].append(part)
            desc = _normalize_text(_cell(row, columns["desc"]))
            if desc and not current_group["descriptions"].get(part):
                current_group["descriptions"][part] = desc

        _finish_group(current_group, parsed_rows)

        return {
            "format": "chenshang",
            "sheet_name": ws.title.strip(),
            "rows": parsed_rows,
            "part_count": len(parsed_rows),
            "manual_split_count": sum(1 for row in parsed_rows if row.get("needs_manual_split")),
        }
    finally:
        workbook.close()


def _normalize_cutoff_for_query(cutoff_date: str) -> str:
    text = str(cutoff_date or "").strip()
    if not text:
        raise ValueError("cutoff_date 為必填")
    if "T" not in text and len(text) == 10:
        return f"{text}T23:59:59.999999"
    return text


def resolve_cutoff_batch(batch_code: str) -> dict:
    """把批次代碼解析成該批最新一次有效發料的 dispatched_at（含該批本身在基準側）。"""
    code = str(batch_code or "").strip()
    if not code:
        raise ValueError("請選擇停損批次")
    for option in db.get_st_reconcile_cutoff_batch_options():
        if option["code"] == code:
            return option
    raise ValueError(f"找不到批次 {code} 的有效發料紀錄（可能已退回），請重新選擇")


def _build_summary() -> dict[str, int]:
    return {
        CATEGORY_HAVE_OURS_NOT_THEIRS: 0,
        CATEGORY_HAVE_THEIRS_NOT_OURS: 0,
        CATEGORY_QTY_MISMATCH: 0,
        CATEGORY_UNATTRIBUTED: 0,
        CATEGORY_MATCHED: 0,
    }


def _build_genlin_summary() -> dict[str, int]:
    return {
        CATEGORY_MATCHED: 0,
        CATEGORY_STOP_LOSS: 0,
        CATEGORY_GENLIN_BLANK_PHYSICAL: 0,
    }


def _classify(diff: float, has_ours_event: bool, tol: float) -> str:
    if abs(diff) <= tol:
        return CATEGORY_MATCHED
    if diff > tol and has_ours_event:
        return CATEGORY_HAVE_OURS_NOT_THEIRS
    if diff < -tol and not has_ours_event:
        return CATEGORY_HAVE_THEIRS_NOT_OURS
    return CATEGORY_UNATTRIBUTED


def _build_genlin_preview(parsed: dict[str, Any], cutoff_date: str, cutoff_for_query: str, tol: float) -> dict[str, Any]:
    source_columns = parsed.get("source_columns") or {"book": "F", "physical": "G"}
    physical_column = source_columns.get("physical") or "G"
    available = set(db.get_st_inventory_stock())
    main_path = db.get_setting('main_file_path')
    if main_path and Path(main_path).is_file():
        available.update(read_stock(main_path))
    parsed = dict(parsed, rows=[dict(row, part_number=_resolve_part(_normalize_part(row['part_number']), available)) for row in parsed['rows']])
    part_numbers = [str(row.get("part_number") or "") for row in parsed["rows"] if row.get("part_number")]
    theoretical = theoretical_stock_with_details(cutoff_for_query, part_numbers=part_numbers)
    stock_by_part = theoretical.get("stock") or {}

    combined: dict[str, dict] = {}
    for row in parsed["rows"]:
        part = str(row.get("part_number") or "").strip().upper()
        if not part:
            continue
        existing = combined.setdefault(part, {
            "part_number": part,
            "description": row.get("description") or "",
            "book_qty": 0.0,
            "physical_qty": 0.0,
            "has_physical": False,
        })
        if row.get("description") and not existing.get("description"):
            existing["description"] = row.get("description")
        existing["book_qty"] += float(row.get("book_qty") or 0)
        if row.get("physical_qty") is not None:
            existing["physical_qty"] += float(row.get("physical_qty") or 0)
            existing["has_physical"] = True

    rows: list[dict] = []
    summary = _build_genlin_summary()
    for part in sorted(combined):
        item = combined[part]
        book_qty = float(item.get("book_qty") or 0)
        theoretical_qty = float(stock_by_part.get(part, 0.0))
        if item.get("has_physical"):
            physical_qty: float | None = float(item.get("physical_qty") or 0)
            book_vs_physical_diff: float | None = round(book_qty - physical_qty, 6)
            diff: float | None = round(physical_qty - theoretical_qty, 6)
            category = CATEGORY_MATCHED if abs(book_vs_physical_diff) <= tol else CATEGORY_STOP_LOSS
            notes: list[str] = []
        else:
            physical_qty = None
            book_vs_physical_diff = None
            diff = None
            category = CATEGORY_GENLIN_BLANK_PHYSICAL
            notes = [f"{physical_column} 欄未填實盤，commit 時不更新 ST 庫存，也不建立停損點基準"]
        summary[category] = int(summary.get(category, 0)) + 1
        rows.append({
            "part_number": part,
            "description": item.get("description") or "",
            "book_qty": book_qty,
            "physical_qty": physical_qty,
            "theoretical": theoretical_qty,
            "book_vs_physical_diff": book_vs_physical_diff,
            "diff": diff,
            "category": category,
            "notes": notes,
        })

    anchor = db.get_latest_st_reconcile_anchor(cutoff_for_query)
    window_start = str(anchor.get("aligned_at") or "") if anchor else ""
    covered_parts = set(combined)
    uncovered_parts = [
        item
        for item in db.get_defective_part_totals(cutoff_for_query, after_at=window_start)
        if item["part_number"] not in covered_parts
    ]

    return {
        "format": "genlin",
        "mode": "stop_loss",
        "cutoff_date": str(cutoff_date or "").strip(),
        "sheet_name": parsed["sheet_name"],
        "source_columns": source_columns,
        "parts": rows,
        "summary": summary,
        "categories": {
            CATEGORY_MATCHED: [row for row in rows if row["category"] == CATEGORY_MATCHED],
            CATEGORY_STOP_LOSS: [row for row in rows if row["category"] == CATEGORY_STOP_LOSS],
            CATEGORY_GENLIN_BLANK_PHYSICAL: [
                row for row in rows if row["category"] == CATEGORY_GENLIN_BLANK_PHYSICAL
            ],
        },
        "uncovered_parts": uncovered_parts,
        "uncovered_window_start": window_start,
        "assumptions": _build_genlin_assumptions(source_columns),
    }


def _batch_preview_token(
    *,
    parsed: dict,
    main_path: str,
    cutoff_at: str,
    batch_code: str,
    session: dict,
    rows: list[dict],
    part_mappings: dict[str, str],
    main_sha256: str,
) -> str:
    """鎖定盤點來源、主檔、工作階段及所有可選料號的試算依據。"""
    payload = {
        "session": {
            key: session.get(key)
            for key in ("id", "cutoff_at", "cutoff_code", "started_at", "status")
        },
        "cutoff_at": cutoff_at,
        "batch_code": batch_code,
        "count_date": parsed.get("count_date") or "",
        "source_sha256": parsed.get("source_sha256") or "",
        "main_sha256": main_sha256,
        "part_mappings": part_mappings,
        "parts": [
            {
                key: row.get(key)
                for key in (
                    "part_number",
                    "source_part_numbers",
                    "book_qty",
                    "physical_qty",
                    "cutoff_main",
                    "current_main",
                    "defect_delta",
                    "defective_record_ids",
                    "expected_count",
                    "target_main",
                    "main_adjustment",
                    "source_rows",
                    "blocked_reason",
                )
            }
            for row in rows
        ],
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _read_reconcile_history(main_path: str, batch_code: str, count_date: str) -> tuple[dict, set[str]]:
    """主檔即調整帳本，還原備份或換年度後只採用檔案仍存在的事件。"""
    workbook = openpyxl.load_workbook(main_path, keep_vba=Path(main_path).suffix.lower() == '.xlsm')
    try:
        ws = workbook.worksheets[0]
        anchor = validate_insertion_anchor(ws, batch_code)
        cutoff_end = max(col + 2 for col in range(anchor, ws.max_column + 1)
                         if str(ws.cell(1, col).value or '').strip() == batch_code)
        history: dict[str, dict] = {}
        reversal_ids: set[str] = set()
        for event in _stock_events(ws):
            kind = event['kind']
            col = int(event['start_col'])
            header = str(ws.cell(1, col).value or '').strip()
            is_defect = header in {'不良品扣帳', '加工多打扣帳', '不良品回復', '加工多打回復'}
            if kind != 'reconcile' and not is_defect:
                continue
            if is_defect and col <= cutoff_end:
                continue
            match = re.search(r'盤點調整 (\d{4}-\d{2}-\d{2})', header)
            event_date = match.group(1) if match else ''
            for row_idx in range(2, ws.max_row + 1):
                part = _normalize_part(ws.cell(row_idx, PART_COL).value)
                increase = ws.cell(row_idx, col).value
                has_usage = int(event['balance_col']) == col + 2
                decrease = ws.cell(row_idx, col + 1).value if has_usage else 0
                if not part or (increase is None and decrease in (None, 0)):
                    continue
                item = history.setdefault(part, {'prior_adjustment': 0.0, 'defect_delta': 0.0, 'blocked_reason': ''})
                if any(value is not None and not isinstance(value, (int, float)) for value in (increase, decrease)):
                    item['blocked_reason'] = '主檔既有盤點或不良品／多打數量無法辨識，請先核對'
                    continue
                delta = float(increase or 0) - float(decrease or 0)
                if is_defect:
                    # 舊版兩欄扣帳以正數表示扣除；三欄則保留加回與扣除的正負淨額。
                    if not has_usage and kind == 'deduct':
                        delta = -delta
                    item['defect_delta'] += delta
                elif not event_date:
                    item['blocked_reason'] = '主檔既有盤點調整缺少日期，請先核對'
                elif event_date > count_date:
                    item['blocked_reason'] = f'此料已套用 {event_date} 盤點，不可再套用較舊的 {count_date} 盤點'
                elif col > cutoff_end:
                    item['prior_adjustment'] += delta
        return history, reversal_ids
    finally:
        workbook.close()


def _build_batch_genlin_preview(parsed: dict, cutoff_at: str, batch_code: str, source_filename: str, tol: float,
                                part_mappings: dict[str, str]) -> dict:
    count_date = parsed.get('count_date') or ''
    if not count_date and source_filename:
        for match in re.finditer(r'(?<!\d)(20\d{2})[/-](\d{1,2})[/-](\d{1,2})(?!\d)', source_filename):
            try:
                count_date = date(*(int(value) for value in match.groups())).isoformat()
                break
            except ValueError:
                continue
    if not count_date:
        raise ValueError('找不到盤點日期，請在盤點表標題或檔名加入 YYYY-MM-DD 日期')
    parsed = dict(parsed, count_date=count_date)
    count_at = count_date + 'T23:59:59.999999'
    if count_at < cutoff_at:
        raise ValueError('盤點日期不可早於截止批次')
    main_path = db.get_setting('main_file_path')
    if not main_path or not Path(main_path).is_file():
        raise ValueError('找不到目前主檔，無法取得截止批次結存')
    session = db.get_active_inventory_count_session('st')
    if not session or session['cutoff_at'] != cutoff_at or session['cutoff_code'] != batch_code:
        raise ValueError('請先開始相同截止批次的盤點並鎖定庫存')
    main_sha256 = _sha256_file(main_path)
    stocks = read_batch_stock(main_path, batch_code)
    history, reversal_ids = _read_reconcile_history(main_path, batch_code, count_date)
    current_st = db.get_st_inventory_stock()
    available = set(read_stock(main_path))
    source_parts = {_normalize_part(row['part_number']) for row in parsed['rows']}
    resolved_parts = {source: _resolve_part(source, available) for source in source_parts}
    for source, target in part_mappings.items():
        if source not in source_parts:
            raise PartMappingError(f'原料號 {source} 不在本次盤點表中')
        if resolved_parts[source] in available:
            raise PartMappingError(f'原料號 {source} 已匹配主檔，不可改成其他料號；若無有效結存請先檢查主檔')
        if target not in stocks:
            raise PartMappingError(f'選定料號 {target} 不是具有有效截止批次及目前結存的主檔料號')
        resolved_parts[source] = target
    sources_by_target: dict[str, list[str]] = defaultdict(list)
    for source, target in resolved_parts.items():
        sources_by_target[target].append(source)
    for target, sources in sources_by_target.items():
        if len(sources) > 1 and any(source in part_mappings for source in sources):
            raise PartMappingError(f'原料號 {"、".join(sorted(sources))} 同時對應到 {target}，請分開核對，不會自動合併')
    candidate_stock = {part: values['current_main'] for part, values in stocks.items() if part not in sources_by_target}
    defects = db.get_defective_interval_parts(cutoff_at, '9999-12-31T23:59:59', reversal_ids=reversal_ids)
    combined = {}
    for source_row in parsed['rows']:
        source_part = _normalize_part(source_row['part_number'])
        part = resolved_parts[source_part]
        row = combined.setdefault(part, dict(part_number=part, source_part_numbers=[], source_rows=[], warnings=[], description=source_row.get('description') or '', book_qty=0.0, physical_qty=None))
        row['source_rows'].append(dict(part_number=source_part, book_qty=source_row.get('book_qty'), physical_qty=source_row.get('physical_qty')))
        if source_part not in row['source_part_numbers']:
            row['source_part_numbers'].append(source_part)
        row['book_qty'] += float(source_row.get('book_qty') or 0)
        if source_row.get('physical_qty') is not None:
            row['physical_qty'] = float(row['physical_qty'] or 0) + float(source_row['physical_qty'])
    rows = []
    uncovered = []
    summary = _build_genlin_summary()
    for part, row in sorted(combined.items()):
        if len(row['source_rows']) > 1:
            row['warnings'].append('多筆來源合併，請逐筆核對料號與實盤數量')
        if any(source['physical_qty'] is None for source in row['source_rows']):
            row['physical_qty'] = None
            row['blocked_reason'] = '來源有未填實盤數量，請補齊後重新試算'
            row['warnings'].append(row['blocked_reason'])
        if part not in stocks:
            can_map = part not in available
            reason = (f'主檔找不到料號 {part}，可人工選擇正確料號後重新試算' if can_map
                      else f'主檔料號 {part} 已存在，但無有效截止批次或目前結存；請先檢查主檔，不可改成其他料號')
            uncovered.append(dict(row,
                                  source_part_number=row['source_part_numbers'][0],
                                  can_map=can_map,
                                  suggestions=suggest_main_part_numbers(part, candidate_stock) if can_map else [],
                                  reason=reason))
            continue
        row.update(stocks[part])
        part_history = history.get(part, {'prior_adjustment': 0.0})
        row['prior_adjustment'] = part_history['prior_adjustment']
        if part_history.get('blocked_reason'):
            row['blocked_reason'] = part_history['blocked_reason']
        if row.get('blocked_reason'):
            row['warnings'].append(row['blocked_reason'])
            row['physical_qty'] = None
        row['defect_delta'] = part_history.get('defect_delta', 0.0)
        row['defective_record_ids'] = defects.get(part, {}).get('defective_record_ids', []) if 'defect_delta' in part_history else []
        row['expected_count'] = round(row['cutoff_main'] + row['defect_delta'] + row['prior_adjustment'], 6)
        row['theoretical'] = row['expected_count']
        row['preserved_delta'] = round(row['current_main'] - row['expected_count'], 6)
        physical = row['physical_qty']
        row['diff'] = round(physical - row['expected_count'], 6) if physical is not None else None
        row['target_main'] = round(physical + row['preserved_delta'], 6) if physical is not None else None
        row['target_current'] = row['target_main']
        row['main_adjustment'] = round(row['target_main'] - row['current_main'], 6) if physical is not None else None
        row['current_st'] = float(current_st.get(part, 0.0))
        row['book_vs_physical_diff'] = round(row['book_qty'] - physical, 6) if physical is not None else None
        row['category'] = CATEGORY_GENLIN_BLANK_PHYSICAL if physical is None else CATEGORY_MATCHED if abs(row['diff']) <= tol else CATEGORY_STOP_LOSS
        row['notes'] = ['後續批次異動保留，現在庫存以主檔結存重算']
        if any(source in part_mappings for source in row['source_part_numbers']):
            row['manual_mapping'] = True
            row['warnings'].extend(f'人工選定料號：{source} → {part}，請重新核對數量'
                                   for source in row['source_part_numbers'] if source in part_mappings)
        summary[row['category']] += 1
        rows.append(row)
    if db.get_setting('main_file_path') != main_path or _sha256_file(main_path) != main_sha256:
        raise ValueError('試算期間主檔已變更，請重新試算')
    preview_token = _batch_preview_token(
        parsed=parsed,
        main_path=main_path,
        cutoff_at=cutoff_at,
        batch_code=batch_code,
        session=session,
        rows=rows,
        part_mappings=part_mappings,
        main_sha256=main_sha256,
    )
    return dict(format='genlin', mode='stop_loss', cutoff_date=cutoff_at, cutoff_batch_code=batch_code,
                count_date=count_date, sheet_name=parsed['sheet_name'], source_columns=parsed['source_columns'],
                preview_token=preview_token, part_mappings=part_mappings, main_sha256=main_sha256,
                parts=rows, summary=summary, categories={key: [row for row in rows if row['category'] == key] for key in summary},
                uncovered_parts=uncovered, assumptions=[
                    _build_genlin_assumptions(parsed['source_columns'])[0],
                    f'盤點日期 {count_date}；截止批次 {batch_code} 後、套用前已存在的不良品／多打淨額併入盤點，不再重複加扣。',
                    '後續發料、用料及其他異動照常計算；只更新勾選料號，原始不良品／多打明細保留查帳。',
                    'ST 庫存僅供警示，不會由本次主檔盤點調整寫入或吸收歷史。',
                    '提交時會重讀來源、主檔與不良明細；資料有變動時必須重新試算。',
                ])


def build_st_reconcile_preview(path: str, cutoff_date: str, *, tol: float = 1e-6,
                               cutoff_batch_code: str = '', source_filename: str = '',
                               part_mappings: dict[str, str] | None = None) -> dict[str, Any]:
    normalized_mappings = normalize_reconcile_part_mappings(part_mappings)
    parsed = parse_st_reconcile_file(path)
    if normalized_mappings and (parsed.get('format') != 'genlin' or not cutoff_batch_code):
        raise PartMappingError('人工料號對應只支援選擇截止批次的庚霖主檔盤點')
    part_numbers = [str(row.get("part_number") or "") for row in parsed["rows"] if row.get("part_number")]
    cutoff_for_query = _normalize_cutoff_for_query(cutoff_date)
    if parsed.get("format") == "genlin":
        if cutoff_batch_code:
            return _build_batch_genlin_preview(parsed, cutoff_for_query, cutoff_batch_code, source_filename, tol, normalized_mappings)
        return _build_genlin_preview(parsed, cutoff_date, cutoff_for_query, tol)

    theoretical = theoretical_stock_with_details(cutoff_for_query, part_numbers=part_numbers)
    stock_by_part = theoretical.get("stock") or {}
    details_by_part = theoretical.get("order_details") or {}

    combined: dict[str, dict] = {}
    for row in parsed["rows"]:
        part = str(row.get("part_number") or "").strip().upper()
        if not part:
            continue
        existing = combined.setdefault(part, {
            "part_number": part,
            "description": row.get("description") or "",
            "physical": 0.0,
            "group_physical": 0.0,
            "needs_manual_split": False,
            "manual_split_notes": [],
        })
        if row.get("description") and not existing.get("description"):
            existing["description"] = row.get("description")
        if row.get("needs_manual_split"):
            existing["needs_manual_split"] = True
            existing["group_physical"] += float(row.get("group_physical") or 0)
            existing["manual_split_notes"].append(
                f"{row.get('customer_code') or '未填客戶編號'} 群組共 {row.get('group_part_count')} 個料號，總盤點數 {row.get('group_physical') or 0:g}"
            )
        else:
            existing["physical"] += float(row.get("physical") or 0)

    rows: list[dict] = []
    summary = _build_summary()
    for part in sorted(combined):
        item = combined[part]
        theoretical_qty = float(stock_by_part.get(part, 0.0))
        notes: list[str] = []
        if item.get("needs_manual_split"):
            category = CATEGORY_UNATTRIBUTED
            physical: float | None = None
            diff: float | None = None
            notes.append("群組多料號需人工拆分")
            notes.extend(item.get("manual_split_notes") or [])
        else:
            physical = float(item.get("physical") or 0)
            diff = round(physical - theoretical_qty, 6)
            has_ours_event = bool(details_by_part.get(part))
            category = _classify(diff, has_ours_event, tol)
            if has_ours_event:
                notes.append(f"截止日前有效 ST 領用事件 {len(details_by_part.get(part) or [])} 筆")
            if category == CATEGORY_QTY_MISMATCH:
                notes.append("盤點表未提供單號，本版不做單對單數量比對")
        summary[category] = int(summary.get(category, 0)) + 1
        rows.append({
            "part_number": part,
            "description": item.get("description") or "",
            "physical": physical,
            "theoretical": theoretical_qty,
            "diff": diff,
            "category": category,
            "notes": notes,
        })

    categories = defaultdict(list)
    for row in rows:
        categories[row["category"]].append(row)

    return {
        "format": parsed.get("format") or "chenshang",
        "mode": "attribution",
        "cutoff_date": str(cutoff_date or "").strip(),
        "sheet_name": parsed["sheet_name"],
        "parts": rows,
        "summary": summary,
        "categories": dict(categories),
        "assumptions": ASSUMPTIONS,
    }


def _copy_adjustment_group_style(ws, start_col: int) -> None:
    """以截止批次最後三欄為樣板，只複製樣式，不複製任何料號資料。"""
    for offset in range(3):
        source_col = start_col - 3 + offset
        target_col = start_col + offset
        source_letter = get_column_letter(source_col)
        target_letter = get_column_letter(target_col)
        source_dimension = ws.column_dimensions[source_letter]
        target_dimension = ws.column_dimensions[target_letter]
        for attr in ("width", "hidden", "bestFit", "outlineLevel", "collapsed"):
            try:
                setattr(target_dimension, attr, getattr(source_dimension, attr))
            except (AttributeError, TypeError):
                pass
        for row_idx in range(1, ws.max_row + 1):
            source = ws.cell(row=row_idx, column=source_col)
            target = ws.cell(row=row_idx, column=target_col)
            target._style = copy(source._style)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)


def _validate_reconcile_rows(ws, insertion_col: int, selected_parts: list[str]) -> dict[str, int]:
    part_rows: dict[str, list[int]] = defaultdict(list)
    for row_idx in range(2, ws.max_row + 1):
        part = _normalize_part(ws.cell(row=row_idx, column=PART_COL).value)
        if part:
            part_rows[part].append(row_idx)

    resolved: dict[str, int] = {}
    events = _stock_events(ws)
    known_columns = {
        col
        for event in events
        for col in range(int(event["start_col"]), int(event["balance_col"]) + 1)
    }
    for part in selected_parts:
        matches = part_rows.get(part) or []
        if not matches:
            raise ValueError(f"主檔找不到料號 {part}，未寫入任何資料")
        if len(matches) != 1:
            raise ValueError(f"主檔料號 {part} 有重複列，無法安全寫入")
        row_idx = matches[0]
        for col in range(insertion_col, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col)
            if col not in known_columns and (
                isinstance(cell.value, (int, float)) or cell.data_type == "f"
            ):
                raise ValueError(f"主檔料號 {part} 後續有無法辨識的庫存欄位，未寫入任何資料")
        for event in events:
            if int(event["start_col"]) < insertion_col:
                continue
            for col in range(int(event["start_col"]), int(event["balance_col"]) + 1):
                if ws.cell(row=row_idx, column=col).data_type == "f":
                    raise ValueError(f"主檔料號 {part} 後續庫存事件含公式，未寫入任何資料")
        resolved[part] = row_idx
    return resolved


def _restore_main_from_backup(backup_path: str, main_path: str) -> None:
    target = Path(main_path)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}-restore-",
        suffix=target.suffix,
        dir=str(target.parent),
    )
    os.close(fd)
    try:
        shutil.copy2(backup_path, temp_name)
        os.replace(temp_name, target)
    finally:
        temp_path = Path(temp_name)
        if temp_path.exists():
            temp_path.unlink()


@serialized_main_file_write
def _commit_batch_genlin(
    path: str,
    cutoff_at: str,
    source_filename: str,
    batch_code: str,
    selected_parts: list[str],
    preview_token: str,
    part_mappings: dict[str, str],
) -> dict:
    """把批次盤點差異寫入主檔；ST 庫存與 ST 歷史完全不變。"""
    try:
        preview = build_st_reconcile_preview(
            path,
            cutoff_at,
            cutoff_batch_code=batch_code,
            source_filename=source_filename,
            part_mappings=part_mappings,
        )
    except PartMappingError:
        raise
    except ValueError as error:
        raise StaleReconcilePreviewError(
            f"盤點來源或主檔已變更，請重新試算後再按確認：{error}"
        ) from error
    if not preview_token or not hmac.compare_digest(preview_token, preview.get("preview_token") or ""):
        raise StaleReconcilePreviewError("盤點來源或主檔已變更，請重新試算後再按確認")

    rows = {row['part_number']: row for row in preview['parts']}
    invalid_parts = sorted(set(selected_parts) - set(rows))
    if invalid_parts:
        raise ValueError(f"勾選的料號不可對帳或不在本次盤點資料中：{'、'.join(invalid_parts)}")
    chosen = [rows[part] for part in selected_parts if rows[part]['physical_qty'] is not None]
    if len(chosen) != len(selected_parts):
        raise ValueError('勾選料號包含未填實盤數量，未寫入任何資料')

    session = db.get_active_inventory_count_session('st')
    if not session or session['cutoff_at'] != cutoff_at or session['cutoff_code'] != batch_code:
        raise ValueError('請先開始相同截止批次的盤點並鎖定庫存')
    main_path = str(db.get_setting('main_file_path') or '').strip()
    if not main_path or not Path(main_path).is_file():
        raise ValueError('找不到目前主檔，未寫入任何資料')
    suffix = Path(main_path).suffix.lower()
    if suffix not in {'.xlsx', '.xlsm'}:
        raise ValueError('目前主檔格式不支援安全寫入，請先轉成 xlsx 或 xlsm')

    workbook = openpyxl.load_workbook(main_path, keep_vba=(suffix == '.xlsm'))
    backup_path = ''
    write_started = False
    snapshot_state: dict | None = None
    try:
        if db.get_setting('main_file_path') != main_path or _sha256_file(main_path) != preview['main_sha256']:
            raise StaleReconcilePreviewError('試算後主檔已變更，請重新試算')
        ws = workbook.worksheets[0]
        anchor_col = validate_insertion_anchor(ws, batch_code)
        matching_cols = [
            col
            for col in range(1, ws.max_column + 1)
            if str(ws.cell(row=1, column=col).value or '').strip() == batch_code
        ]
        insertion_col = anchor_col + 3 * len(matching_cols)
        selected_rows = _validate_reconcile_rows(ws, insertion_col, selected_parts)

        insert_columns(ws, insertion_col, 3)
        _copy_adjustment_group_style(ws, insertion_col)
        ws.cell(row=1, column=insertion_col).value = f"盤點調整 {preview['count_date']} 增加"
        ws.cell(row=1, column=insertion_col + 1).value = '扣除數量'
        ws.cell(row=1, column=insertion_col + 2).value = '結存'
        from .inventory_restore_guard import COUNT_COMMENT_PREFIX
        records = db.get_defective_records()
        boundary = dict(parts=[row['part_number'] for row in chosen], cutoff_at=cutoff_at,
                        max_record_id=max((int(row['id']) for row in records), default=0))
        ws.cell(row=1, column=insertion_col).comment = Comment(
            COUNT_COMMENT_PREFIX + json.dumps(boundary, ensure_ascii=False), 'shipping-scheduler',
        )

        result_parts: list[dict] = []
        adjustments: list[dict] = []
        for row in chosen:
            part = row['part_number']
            row_idx = selected_rows[part]
            adjustment = round(float(row['main_adjustment']), 6)
            ws.cell(row=row_idx, column=insertion_col).value = _display_number(max(adjustment, 0.0))
            ws.cell(row=row_idx, column=insertion_col + 1).value = _display_number(max(-adjustment, 0.0))
            recalc = recalc_batch_balances_for_cell(ws, row=row_idx, col=insertion_col)
            if not recalc.get('recalculated') or recalc.get('current_stock') is None:
                raise ValueError(f"主檔料號 {part} 無法安全重算後續結存，未寫入任何資料")
            main_after = round(float(recalc['current_stock']), 6)
            target_main = round(float(row['target_main']), 6)
            if abs(main_after - target_main) > 1e-6:
                raise ValueError(
                    f"主檔料號 {part} 重算結果 {main_after:g} 與目標 {target_main:g} 不一致，未寫入任何資料"
                )
            item = dict(
                row,
                main_before=float(row['current_main']),
                main_after=main_after,
                target_main=target_main,
                target_current=target_main,
                main_adjustment=adjustment,
            )
            result_parts.append(item)
            adjustments.append({
                'part_number': part,
                'adjust_qty': adjustment,
                'main_before': float(row['current_main']),
                'main_after': main_after,
            })

        if db.get_setting('main_file_path') != main_path or _sha256_file(main_path) != preview['main_sha256']:
            raise StaleReconcilePreviewError('寫入前主檔已變更，請重新試算')
        snapshot_state = db.capture_inventory_snapshot_state()
        backup_path = backup_main_file(main_path, str(BACKUP_DIR))
        write_started = True
        _save_workbook_atomically(workbook, main_path)
        snapshot_count = refresh_snapshot_from_main(main_path)
        if snapshot_count <= 0:
            raise RuntimeError('主檔已寫入但庫存快照同步失敗')
        summary = dict(
            alignment_id=None,
            session_id=int(session['id']),
            part_count=len(result_parts),
            updated_count=len(result_parts),
            adjusted_count=sum(abs(row['adjust_qty']) > 1e-6 for row in adjustments),
            total_abs_adjust_qty=round(sum(abs(row['adjust_qty']) for row in adjustments), 6),
            absorbed_defective_records=0,
            absorbed_supplements=0,
            snapshot_count=snapshot_count,
        )
        response = dict(
            ok=True,
            format='genlin',
            mode='stop_loss',
            count_date=preview['count_date'],
            cutoff_batch_code=batch_code,
            source_columns=preview['source_columns'],
            preview_token=preview['preview_token'],
            part_mappings=preview['part_mappings'],
            backup_path=backup_path,
            session_completed=True,
            summary=summary,
            preview_summary=preview['summary'],
            parts=result_parts,
            adjustments=adjustments,
        )
        workbook.close()
        workbook = None
        # 工作階段完成是最後一個可能失敗的步驟；成功後直接回傳，不再執行可能拋錯的工作。
        if not db.finish_inventory_count_session(
            int(session['id']),
            status='completed',
            source_filename=source_filename,
        ):
            raise RuntimeError('盤點工作階段完成失敗')
        return response
    except Exception as error:
        if backup_path and write_started:
            rollback_errors: list[str] = []
            try:
                _restore_main_from_backup(backup_path, main_path)
            except Exception as rollback_error:
                rollback_errors.append(f"主檔還原失敗：{rollback_error}")
            try:
                db.restore_inventory_snapshot_state(snapshot_state or {})
                try:
                    from app.routers.main_file import invalidate_main_data_cache
                    invalidate_main_data_cache()
                except ImportError:
                    pass
            except Exception as rollback_error:
                rollback_errors.append(f"快照狀態還原失敗：{rollback_error}")
            if rollback_errors:
                raise RuntimeError(f"{error}；{'；'.join(rollback_errors)}") from error
        raise
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def commit_st_reconcile_stop_loss(
    path: str,
    cutoff_date: str,
    *,
    source_filename: str = "",
    part_numbers: list[str] | None = None,
    cutoff_label: str = "",
    preview_token: str = "",
    part_mappings: dict[str, str] | None = None,
) -> dict[str, Any]:
    normalized_mappings = normalize_reconcile_part_mappings(part_mappings)
    selected_parts = _normalize_part_numbers(part_numbers)
    if part_numbers is not None and not selected_parts:
        raise ValueError("請至少勾選 1 支料號再建立停損點")

    if cutoff_label:
        if not selected_parts:
            raise ValueError("請至少勾選 1 支料號再寫入主檔")
        return _commit_batch_genlin(
            path,
            cutoff_date,
            source_filename or Path(path).name,
            cutoff_label,
            selected_parts,
            str(preview_token or '').strip(),
            normalized_mappings,
        )

    preview = build_st_reconcile_preview(path, cutoff_date, cutoff_batch_code=cutoff_label,
                                        source_filename=source_filename, part_mappings=normalized_mappings)
    if preview.get("format") != "genlin":
        raise ValueError("停損點 commit 目前只支援庚霖實際庫存格式")
    selected_part_set = set(selected_parts or [])

    cutoff_for_anchor = _normalize_cutoff_for_query(cutoff_date)
    current_stock = db.get_st_inventory_stock()
    audit_delta_rows = db.get_st_inventory_audit_deltas(
        "9999-12-31T23:59:59",
        after_at=cutoff_for_anchor,
        part_numbers=selected_parts or None,
        exclude_reason=ST_RECONCILE_ADJUSTMENT_REASON,
    )
    post_cutoff_delta_by_part: dict[str, float] = defaultdict(float)
    for audit_row in audit_delta_rows:
        audit_part = str(audit_row.get("part_number") or "").strip().upper()
        if not audit_part:
            continue
        post_cutoff_delta_by_part[audit_part] += float(audit_row.get("delta") or 0)
    stock_updates: dict[str, float] = {}
    alignment_parts: list[dict] = []
    adjustments: list[dict] = []

    for row in preview.get("parts") or []:
        part = str(row.get("part_number") or "").strip().upper()
        if not part:
            continue
        if selected_parts is not None and part not in selected_part_set:
            continue
        if row.get("physical_qty") is None:
            continue
        aligned_qty = float(row.get("physical_qty") or 0)
        current_qty = float(current_stock.get(part, 0.0))
        theoretical_qty = float(row.get("theoretical") or 0)
        target_qty = round(aligned_qty + post_cutoff_delta_by_part.get(part, 0.0), 6)
        adjust_qty = round(target_qty - current_qty, 6)
        stock_updates[part] = target_qty
        alignment_parts.append({
            "part_number": part,
            "theoretical_qty": theoretical_qty,
            "physical_qty": aligned_qty,
            "diff": float(row.get("diff") or 0),
            "category": str(row.get("category") or ""),
            "aligned_qty": aligned_qty,
        })
        adjustments.append({
            "part_number": part,
            "adjust_qty": adjust_qty,
            "reason": ST_RECONCILE_ADJUSTMENT_REASON,
            "actor": "reconcile",
        })

    updated_count = db.update_st_inventory_stock(
        stock_updates,
        reason=ST_RECONCILE_ADJUSTMENT_REASON,
        actor="reconcile",
    )
    alignment_id = db.create_st_reconcile_alignment(
        aligned_at=cutoff_for_anchor,
        source_filename=source_filename or Path(path).name,
        note=(
            f"停損點模式：以庚霖實盤 {preview['source_columns']['physical']} 欄重設 ST 庫存基準（截止批次 {cutoff_label}）"
            if cutoff_label
            else f"停損點模式：以庚霖實盤 {preview['source_columns']['physical']} 欄重設 ST 庫存基準"
        ),
        parts=alignment_parts,
        adjustments=adjustments,
    )
    absorbed = db.mark_inventory_history_absorbed(
        alignment_id,
        cutoff_for_anchor,
        [row["part_number"] for row in alignment_parts],
    )
    summary = {
        "alignment_id": alignment_id,
        "aligned_at": cutoff_for_anchor,
        "part_count": len(alignment_parts),
        "updated_count": updated_count,
        "adjusted_count": sum(1 for row in adjustments if abs(float(row.get("adjust_qty") or 0)) > 1e-6),
        "total_abs_adjust_qty": round(sum(abs(float(row.get("adjust_qty") or 0)) for row in adjustments), 6),
        "absorbed_defective_records": absorbed["defective_records"],
        "absorbed_supplements": absorbed["supplements"],
    }
    return {
        "ok": True,
        "format": "genlin",
        "mode": "stop_loss",
        "source_columns": preview.get("source_columns") or {"book": "F", "physical": "G"},
        "summary": summary,
        "preview_summary": preview.get("summary") or {},
        "parts": alignment_parts,
        "adjustments": adjustments,
    }
