import tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from openpyxl.styles import PatternFill
from app.services.merge_to_main import merge_order_batches_to_main,preview_order_batches,restore_main_from_backup_reference
from app.services.main_reconcile import verify_main_write

class MainInsertionTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.path=Path(self.tmp.name)/'main.xlsx';self.backup=Path(self.tmp.name)/'backups'
  wb=Workbook();ws=wb.active
  ws.append(['料號','廠商','MOQ','說明','','','','結存','10-0','PO','MODEL','不良品扣帳','使用數量','結存','11-0','PO','MODEL'])
  ws.append(['PART-A','V',1,'A',None,None,None,100,None,20,80,None,3,77,None,5,72])
  ws['J2'].fill=PatternFill('solid',fgColor='00FF00');ws.column_dimensions['J'].width=23
  sheet=wb.create_sheet('ref');sheet['A1']='=Sheet!$K$2'
  wb.save(self.path);wb.close()
 def batch(self,oid=1,qty=10,groups=1):
  return dict(order_id=oid,model='MODEL',insert_before_code='10-0',groups=[dict(batch_code='10-0',po_number='NEW',bom_model='MODEL',components=[dict(part_number='PART-A',needed_qty=qty,prev_qty_cs=0)]) for _ in range(groups)])
 def commit(self,batches):return merge_order_batches_to_main(str(self.path),batches,backup_dir=str(self.backup))
 def test_insert_keeps_adjustment_other_group_and_format(self):
  result=self.commit([self.batch()]);wb=load_workbook(self.path);ws=wb.active
  self.assertEqual([ws.cell(2,c).value for c in (11,14,17,20)],[90,70,67,62])
  self.assertEqual([ws.cell(1,c).value for c in (9,12,15,18)],['10-0','10-1','不良品扣帳','11-0'])
  self.assertEqual(ws['M2'].fill.fgColor.rgb,'0000FF00');self.assertEqual(ws.column_dimensions['M'].width,23)
  self.assertEqual(wb['ref']['A1'].value,'=Sheet!$N$2');wb.close()
  self.assertTrue(verify_main_write(str(self.path), result['plan_rows'])['ok'])
 def test_preview_reports_downstream_shortage_without_writing(self):
  before=self.path.read_bytes();preview=preview_order_batches(str(self.path),[self.batch(qty=90)],{})
  self.assertEqual(self.path.read_bytes(),before)
  self.assertEqual(max(x['shortage_amount'] for x in preview['shortages']),18)
 def test_multiple_bom_and_replay(self):
  result=self.commit([self.batch(groups=2),self.batch(oid=2)])
  self.assertTrue(verify_main_write(str(self.path), result['plan_rows'])['ok'])
  wb=load_workbook(self.path);ws=wb.active
  self.assertEqual([ws.cell(2,c).value for c in (11,14,17,20,23,26)],[90,80,70,50,47,42]);wb.close()
  restore_main_from_backup_reference(result['backup_path'],self.path,before_order_id=2)
  wb=load_workbook(self.path);ws=wb.active
  self.assertEqual([ws.cell(2,c).value for c in (11,14,17,20,23)],[90,80,60,57,52]);wb.close()
 def test_continuous_insert(self):
  self.commit([self.batch()]);self.commit([self.batch(oid=2)])
  wb=load_workbook(self.path);ws=wb.active
  self.assertEqual([ws.cell(1,c).value for c in (9,12,15)],['10-0','10-1','10-2'])
  self.assertEqual(ws.cell(2,23).value,52);wb.close()

 def test_unknown_history_rejected_without_changing_main(self):
  wb=load_workbook(self.path);wb.active.cell(1,15).value=None;wb.save(self.path);wb.close()
  before=self.path.read_bytes()
  with self.assertRaisesRegex(ValueError,'無法辨識'):
   self.commit([self.batch()])
  self.assertEqual(self.path.read_bytes(),before)
 def test_manifest_failure_restores_original(self):
  before=self.path.read_bytes()
  with patch('app.services.merge_to_main._write_json_atomically',side_effect=OSError('failed')):
   with self.assertRaises(OSError):self.commit([self.batch()])
  self.assertEqual(self.path.read_bytes(),before)
 def test_rollback_checks_codes_before_touching_workbook(self):
  from app.services.dispatch_pipeline import rollback_dispatch_sessions
  session=dict(id=1,order_id=1,backup_path='backup.json',main_file_path=str(self.path))
  with patch('app.services.dispatch_pipeline.get_dispatch_rollback_unavailable_reason',return_value=''), \
       patch('app.services.dispatch_pipeline.db.get_setting',return_value=str(self.path)), \
       patch('app.services.dispatch_pipeline.db.resolve_managed_path',side_effect=lambda p,k:p), \
       patch('app.services.dispatch_pipeline.read_dispatch_code_changes',return_value=[[{'id':1}]]), \
       patch('app.services.dispatch_pipeline.db.validate_dispatch_code_shift_restores',side_effect=ValueError('changed')), \
       patch('app.services.dispatch_pipeline.restore_main_from_backup_reference') as restore:
   with self.assertRaises(ValueError):rollback_dispatch_sessions([session])
   restore.assert_not_called()

if __name__=='__main__':unittest.main()
