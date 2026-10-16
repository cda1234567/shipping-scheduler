from __future__ import annotations

import os
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.models import BomEditorComponentUpdate, BomEditorSaveRequest
from app.services.bom_editor import apply_bom_editor_changes, read_bom_editor_scrap_factors
from app.routers import bom as bom_router
from app.services import merge_drafts
from app.services.bom_quantity import get_component_effective_needed_qty
from app.services.bom_parser import parse_bom
from app.services.merge_drafts import _write_draft_files
from main import app


class BomScrapEditorTests(unittest.TestCase):
    def _workbook(self, path):
        workbook = Workbook()
        ws = workbook.active
        ws["H1"] = 123
        ws["K1"] = 100
        ws["G2"] = 100
        ws["C2"] = "MODEL-A"
        ws["E3"] = "拋料率"
        for row, part, rate, formula in [
            (5, "PART-Z", 0.006, "=ROUNDUP(B5*$K$1*(1+E5),0)"),
            (6, "PART-A", "=2.5%", "=B6*$K$1*(1+E6)"),
            (8, "PART-C", 0.006, "=(G$2*B8)+(G$2*B8)*E8+1"),
            (9, "PART-D", 0.006, 100.6),
        ]:
            ws.cell(row, 2, 1)
            ws.cell(row, 3, part)
            ws.cell(row, 5, rate)
            ws.cell(row, 6, formula)
            ws.cell(row, 9, f"=SUM(G{row},H{row})")
            ws.cell(row, 10, f"=I{row}-F{row}")
        ws["A7"] = "客供"
        workbook.save(path)
        workbook.close()

    def _request(self, rates):
        return BomEditorSaveRequest(po_number=123, order_qty=100, model="MODEL-A", components=[
            BomEditorComponentUpdate(source_row=row, part_number=part, qty_per_board=1, **({"scrap_factor": rates[row]} if row in rates else {}))
            for row, part in [(5, "PART-Z"), (6, "PART-A"), (8, "PART-C"), (9, "PART-D")]
        ])

    def test_changed_rates_round_trip_preserve_other_rates_and_formula_adjustments(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            apply_bom_editor_changes(str(path), self._request({5: 0, 8: 0.02, 9: 0.1}))
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["E5"].value, 0)
            self.assertEqual(workbook.active["E6"].value, "=2.5%")
            self.assertEqual(workbook.active["F5"].value, "=ROUNDUP(B5*$K$1*(1+E5),0)")
            self.assertEqual(workbook.active["F8"].value, "=(G$2*B8)+(G$2*B8)*E8+1")
            self.assertEqual(workbook.active["J8"].value, "=I8-F8")
            self.assertTrue(workbook.calculation.fullCalcOnLoad)
            workbook.close()
            parsed = parse_bom(str(path), "bom-1", path.name, "")
            self.assertEqual(read_bom_editor_scrap_factors(str(path)), {5: 0, 6: 0.025, 7: 0, 8: 0.02, 9: 0.1})
            # 排程沿用既有 F 反推有效率；編輯頁顯示的實際率不改變此政策。
            self.assertAlmostEqual(parsed.components[2].scrap_factor, 0.03)
            for component, expected in zip(parsed.components, (100, 102.5, 103, 110)):
                self.assertAlmostEqual(component.needed_qty, expected)

            record = {"id": "bom-1", "filepath": str(path), "filename": path.name}
            with patch("app.services.merge_drafts.db.get_bom_file", return_value=record), \
                 patch("app.services.merge_drafts._ensure_editable_bom_for_draft", return_value=record), \
                 patch("app.services.merge_drafts.save_workbook_with_recalc", side_effect=lambda wb, target: wb.save(target)):
                files = _write_draft_files(1, [{"bom_file_id": "bom-1", "order_qty": 100}], root_dir=Path(directory) / "drafts")
            generated = parse_bom(files[0]["filepath"], "draft", "draft.xlsx", "")
            self.assertEqual([component.part_number for component in generated.components], ["PART-A", "PART-Z", "PART-C", "PART-D"])
            self.assertEqual({c.part_number: (c.scrap_factor, c.needed_qty) for c in generated.components}, {c.part_number: (c.scrap_factor, c.needed_qty) for c in parsed.components})

    def test_hardcoded_rate_updates_without_dropping_roundup(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            workbook = load_workbook(path)
            workbook.active["F5"] = "=ROUNDUP(B5*$K$1*1.06,0)"
            workbook.active["E5"] = 0.06
            workbook.save(path)
            workbook.close()
            apply_bom_editor_changes(str(path), self._request({5: 0.015}))
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["F5"].value, "=ROUNDUP(B5*$K$1*(1+E5),0)")
            workbook.close()
            first = parse_bom(str(path), "bom-1", path.name, "").components[0]
            self.assertEqual(read_bom_editor_scrap_factors(str(path))[5], 0.015)
            self.assertEqual(first.needed_qty, 102)

    def test_api_save_backups_revision_reload_and_unsupported_formula_rollback(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"PYTEST_CURRENT_TEST": "1"}):
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            record = {"id": "bom-1", "filepath": str(path), "filename": path.name, "uploaded_at": "2026-10-15"}
            client = TestClient(app)
            with patch("app.routers.bom._get_required_bom", return_value=record), \
                 patch("app.routers.bom.ensure_bom_revision_history"), \
                 patch("app.services.bom_editor._BOM_BACKUP_DIR", Path(directory)), \
                 patch("app.routers.bom.db.save_bom_file") as save, \
                 patch("app.routers.bom.snapshot_bom_revision") as snapshot, \
                 patch("app.routers.bom.db.log_activity"):
                response = client.put("/api/bom/bom-1/editor", json=self._request({5: 0.015, 8: 0.02}).dict())
                self.assertEqual(response.status_code, 200, response.text)
                self.assertTrue(Path(response.json()["backup_path"]).exists())
                snapshot.assert_called_once()
                saved = client.get("/api/bom/bom-1/editor").json()
                self.assertEqual(saved["components"][0]["scrap_factor"], 0.015)
                self.assertEqual(saved["components"][0]["needed_qty"], 102)
                self.assertEqual(saved["components"][1]["scrap_factor"], 0.025)
                self.assertEqual(saved["components"][2]["scrap_factor"], 0.02)
                self.assertAlmostEqual(save.call_args.args[0]["components"][2]["scrap_factor"], 0.03)
                workbook = load_workbook(path)
                workbook.active["F5"] = '=IF(B5>0,105,100)'
                workbook.save(path)
                workbook.close()
                original = path.read_bytes()
                save.reset_mock()
                response = client.put("/api/bom/bom-1/editor", json=self._request({5: 0}).dict())
                self.assertEqual(response.status_code, 400)
                self.assertIn("無法辨識拋料率", response.json()["detail"])
                self.assertEqual(path.read_bytes(), original)
                save.assert_not_called()

    def test_rejects_extra_term_constant_without_changing_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            workbook = load_workbook(path)
            workbook.active["F5"] = "=B5*G$2+2*0.006"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "無法辨識拋料率"):
                apply_bom_editor_changes(str(path), self._request({5: 0.02}))
            self.assertEqual(path.read_bytes(), original)

    def test_numeric_needed_quantity_keeps_changed_rate_in_both_export_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            apply_bom_editor_changes(str(path), self._request({9: 0.02}))
            parsed = parse_bom(str(path), "bom-1", path.name, "")
            component = parsed.components[-1].dict()
            for export in (bom_router, merge_drafts):
                for target_quantity in (100, 200):
                    with self.subTest(export=export.__name__, target=target_quantity):
                        workbook = load_workbook(path)
                        export._apply_target_order_qty_to_ws(workbook.active, target_quantity, 100)
                        self.assertAlmostEqual(workbook.active["F9"].value, get_component_effective_needed_qty(component, target_quantity, 100))
                        workbook.close()

    def test_editor_reads_explicit_zero_even_with_another_rate_in_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsx"
            self._workbook(path)
            workbook = load_workbook(path)
            workbook.active["E5"] = 0
            workbook.active["L5"] = 0.02
            workbook.active["F5"] = "=B5*G$2*(1+E5)*(1+L5)"
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_bom_editor_scrap_factors(str(path))[5], 0)
            # 不更改既有排程解析對其他公式加項的處理方式。
            self.assertAlmostEqual(parse_bom(str(path), "bom-1", path.name, "").components[0].scrap_factor, 0.02)

    def test_rate_edit_preserves_macro_archive(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bom.xlsm"
            self._workbook(path)
            macro_payload = b"macro-archive-regression-fixture"
            with zipfile.ZipFile(path, "a") as archive:
                archive.writestr("xl/vbaProject.bin", macro_payload)
            apply_bom_editor_changes(str(path), self._request({5: 0.02}))
            with zipfile.ZipFile(path) as archive:
                self.assertEqual(archive.read("xl/vbaProject.bin"), macro_payload)
            workbook = load_workbook(path, keep_vba=True)
            self.assertEqual(workbook.active["E5"].value, 0.02)
            self.assertEqual(workbook.active["F5"].value, "=ROUNDUP(B5*$K$1*(1+E5),0)")
            workbook.close()
            workbook.vba_archive.close()

    def test_rate_validation_rejects_nonfinite_negative_and_100_percent(self):
        for rate in (-0.01, 1, float("nan"), float("inf")):
            with self.subTest(rate=rate), self.assertRaises(ValidationError):
                BomEditorComponentUpdate(source_row=5, part_number="PART-A", scrap_factor=rate)
        self.assertEqual(BomEditorComponentUpdate(source_row=5, part_number="PART-A", scrap_factor=0).scrap_factor, 0)

    def test_frontend_percentage_display_only_sends_changed_rates_and_rejects_invalid_values(self):
        root = Path(__file__).resolve().parents[1]
        script = r'''
const fs = require('fs');
const assert = require('assert');
const text = fs.readFileSync('static/modules/bom_manager.js', 'utf8');
const source = text.slice(text.indexOf('function compareEditorComponents'), text.indexOf('async function saveEditor'));
const fields = Object.fromEntries(['po', 'order-qty', 'model', 'pcb', 'group-model', 'title', 'source', 'source-format', 'convert-note', 'table-body'].map(name => [`bom-editor-${name}`, {value:'', style:{}}]));
let rows = [];
const document = { getElementById: id => fields[id], querySelectorAll: () => rows };
const {renderEditor, readEditorPayload} = new Function('document', 'esc', source + '; return {renderEditor, readEditorPayload};')(document, value => String(value));
renderEditor({filename:'test.xlsx', model:'TEST', order_qty:100, components:[{source_row:5, part_number:'A', qty_per_board:1, scrap_factor:0.006}, {source_row:6, part_number:'B', scrap_factor:0.025}]});
assert(fields['bom-editor-table-body'].innerHTML.includes('value="0.6" data-original="0.6"'));
assert(fields['bom-editor-table-body'].innerHTML.includes('value="2.5" data-original="2.5"'));
function makeRow(row, rate) {
  const inputs = {part_number:{value:`P-${row}`}, description:{value:''}, qty_per_board:{value:'1'}, needed_qty:{value:'100'}, scrap_percent:{value:rate, dataset:{original:rate}}};
  return {dataset:{row:String(row)}, inputs, querySelector: selector => inputs[selector.match(/"([^"]+)"/)[1]]};
}
rows = [makeRow(5, '0.6'), makeRow(6, '2.5')];
assert(!('scrap_factor' in readEditorPayload().components[0]));
rows[0].inputs.scrap_percent.value = '0';
assert.strictEqual(readEditorPayload().components[0].scrap_factor, 0);
assert(!('scrap_factor' in readEditorPayload().components[1]));
rows[0].inputs.scrap_percent.value = '1.5';
assert.strictEqual(readEditorPayload().components[0].scrap_factor, 0.015);
for (const value of ['', '-1', '100', 'NaN', 'Infinity']) {
  rows[0].inputs.scrap_percent.value = value;
  assert.throws(readEditorPayload, /拋料率/);
}
console.log('percentage editing passed');
'''
        completed = subprocess.run(["node", "-e", script], cwd=root, capture_output=True, text=True)
        self.assertEqual(completed.returncode, 0, completed.stderr)
