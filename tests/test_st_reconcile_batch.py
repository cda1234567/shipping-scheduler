from __future__ import annotations

import asyncio
import hashlib
import io
import json
import sqlite3
import unittest
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import database as db
from app.routers import reconcile as reconcile_router
from app.services import st_reconcile
from app.services.main_reader import read_batch_stock, read_stock
from app.services.st_reconcile import (
    StaleReconcilePreviewError,
    build_st_reconcile_preview,
    commit_st_reconcile_stop_loss,
    parse_st_reconcile_file,
)


class BatchReconcileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.conn = sqlite3.connect(':memory:', check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(db._CREATE_SQL)
        self.addCleanup(self.conn.close)

        @contextmanager
        def connection():
            with self.conn:
                yield self.conn

        self.patcher = patch.object(db, 'get_conn', connection)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.backup_patcher = patch.object(st_reconcile, 'BACKUP_DIR', Path(self.tmp.name) / 'backups')
        self.backup_patcher.start()
        self.addCleanup(self.backup_patcher.stop)

        self.count = Path(self.tmp.name) / 'count_2026-9-23.xlsx'
        wb = Workbook()
        ws = wb.active
        ws['A1'] = '盤點日期 2026/9/23'
        for col, value in {4: 'consign invoice NO', 6: '辰尚庫存', 7: '庚霖庫存 當下實際', 21: '實際庫存總和', 22: '生產結餘'}.items():
            ws.cell(3, col, value)
        source_rows = (
            ('EC-20128A-TAB', 780, 771, 33179, 33200),
            ('EC-30037A-TAB', -2848, 351, 99999, 88888),
            ('PART-KEEP', 50, 50, 777, 666),
        )
        for row_idx, values in enumerate(source_rows, start=5):
            for col, value in zip((4, 6, 7, 21, 22), values):
                ws.cell(row_idx, col, value)
        wb.save(self.count)
        wb.close()

        self.main = Path(self.tmp.name) / 'main.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['料號', '廠商', 'MOQ', None, None, None, None, '盤點', '9-11', None, '結存', '9-12', None, '結存', '9-12', None, '結存', '9-13', None, '結存'])
        ws.append(['EC-20128A-TAB', '', 1000, None, None, None, None, 900, 0, 10, 890, 0, 90, 800, 0, 20, 780, 0, 145, 635])
        ws.append(['EC-30037A-TAB', '', 1000, None, None, None, None, -2800, 0, 20, -2820, 0, 28, -2848, 0, 0, -2848, 3617, 0, 769])
        ws.append(['PART-KEEP', '', 1000, None, None, None, None, 100, 0, 10, 90, 0, 20, 70, 0, 20, 50, 0, 0, 50])
        wb.save(self.main)
        wb.close()
        db.set_setting('main_file_path', str(self.main))
        self.conn.execute(
            "INSERT INTO st_inventory_snapshot(part_number,stock_qty) VALUES"
            "('EC-20128A-TAB',22000),('EC-30037A-TAB',9236),('PART-KEEP',444)"
        )
        self.conn.executemany(
            "INSERT INTO inventory_snapshot(part_number,stock_qty,moq,moq_manual,description,snapshot_at) "
            "VALUES(?,?,?,?,?,?)",
            [
                ('EC-20128A-TAB', 635, 111, 1, '原始說明 A', '2026-09-19T09:00:01'),
                ('EC-30037A-TAB', 769, 222, 0, '原始說明 B', '2026-09-19T09:00:02'),
                ('PART-KEEP', 50, 333, 1, '原始說明 C', '2026-09-19T09:00:03'),
            ],
        )
        db.set_setting('main_part_count', 'before-count')
        self.conn.executemany("INSERT INTO defective_records(part_number,defective_qty,status,created_at) VALUES(?,?,?,?)", [
            ('EC-20128A-TAB', 4, 'open', '2026-09-20T10:00:00'),
            ('EC-20128A-TAB', 5, 'confirmed', '2026-09-23T12:00:00'),
            ('EC-20128A-TAB', 100, 'open', '2026-09-24T00:00:00'),
        ])
        db.start_inventory_count_session(cutoff_at='2026-09-19T10:00:00', cutoff_code='9-12')

    def preview(self, mappings=None):
        return build_st_reconcile_preview(
            str(self.count),
            '2026-09-19T10:00:00',
            cutoff_batch_code='9-12',
            source_filename=self.count.name,
            part_mappings=mappings,
        )

    def commit(self, *, parts=None, token=None, mappings=None):
        if token is None:
            token = self.preview(mappings)['preview_token']
        return commit_st_reconcile_stop_loss(
            str(self.count),
            '2026-09-19T10:00:00',
            cutoff_label='9-12',
            source_filename=self.count.name,
            part_numbers=parts or ['EC-20128A-TAB'],
            preview_token=token,
            part_mappings=mappings,
        )

    def set_count_part(self, row, part):
        wb = load_workbook(self.count)
        wb.active.cell(row, 4).value = part
        wb.save(self.count)
        wb.close()

    def invalidate_cutoff(self, row):
        wb = load_workbook(self.main)
        for col in (8, 11, 14, 17):
            wb.active.cell(row, col).value = None
        wb.save(self.main)
        wb.close()

    @staticmethod
    def file_hash(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def prepare_count_example(self, *, defect_supply=0, defect_usage=20, legacy=False):
        wb = Workbook()
        ws = wb.active
        ws.append(['料號', '廠商', 'MOQ', None, None, None, None, '盤點',
                   '9-12', None, '結存', '不良品扣帳', '使用數量', '結存', '10-1', None, '結存'])
        ws.append(['EC-20128A-TAB', '', 1, None, None, None, None, 90,
                   0, 0, 90, defect_supply, defect_usage, 90 + defect_supply - defect_usage,
                   20, 40, 70 + defect_supply - defect_usage])
        if legacy:
            ws.delete_cols(13)
            ws.cell(1, 13).value = '結存'
            ws.cell(2, 12).value = defect_usage
        wb.save(self.main)
        wb.close()
        wb = load_workbook(self.count)
        wb.active['G5'] = 70
        wb.save(self.count)
        wb.close()

    def test_count_absorbs_defect_but_keeps_later_supply_and_usage(self):
        self.prepare_count_example()
        result = self.commit()['parts'][0]
        self.assertEqual(result['defect_delta'], -20)
        self.assertEqual(result['main_after'], 50)  # Andy：70 + 發20 - 用40。
        wb = load_workbook(self.main)
        ws = wb.active
        self.assertEqual([ws.cell(2, c).value for c in (15, 16, 18, 19, 20)], [0, 20, 20, 40, 50])
        wb.close()
        self.restart_session()
        result = self.commit()['parts'][0]
        self.assertEqual(result['main_adjustment'], 0)
        self.assertEqual(result['main_after'], 50)

    def test_count_absorbs_signed_overrun_and_legacy_deduction(self):
        for supply, usage, legacy in ((0, 20, False), (30, 10, False), (0, 20, True)):
            with self.subTest(supply=supply, usage=usage, legacy=legacy):
                self.restart_session()
                self.prepare_count_example(defect_supply=supply, defect_usage=usage, legacy=legacy)
                wb = load_workbook(self.main)
                wb.active.cell(1, 12).value = '加工多打扣帳'
                wb.save(self.main)
                wb.close()
                result = self.commit()['parts'][0]
                self.assertEqual(result['main_after'], 50)
                self.assertEqual(result['defect_delta'], supply - usage)

    def test_count_protects_ambiguous_old_records_but_not_new_or_restored_main(self):
        from app.services.inventory_restore_guard import is_count_protected_record, ensure_defective_batch_delete_allowed
        from app.routers import defectives
        self.prepare_count_example()
        before = self.main.read_bytes()
        old = db.get_defective_records()[0]  # 含盤點檔日期之後才建立的舊扣帳。
        self.commit()
        self.assertTrue(is_count_protected_record(old))
        with self.assertRaisesRegex(HTTPException, '400'):
            ensure_defective_batch_delete_allowed({'items': [old]})
        with self.assertRaises(HTTPException):
            asyncio.run(defectives.delete_record(old['id']))
        new_id = db.create_defective_record(dict(part_number=old['part_number'], defective_qty=1))
        self.assertFalse(is_count_protected_record(db.get_defective_record(new_id)))
        self.main.write_bytes(before)
        self.assertFalse(is_count_protected_record(old))

    def test_newer_count_left_of_cutoff_still_blocks_older_count(self):
        from app.services.main_insertion import insert_columns
        wb = load_workbook(self.main)
        ws = wb.active
        insert_columns(ws, 9, 3)
        for col, value in ((9, '盤點調整 2026-09-24 增加'), (10, '扣除數量'), (11, '結存')):
            ws.cell(1, col).value = value
        for col, value in ((9, 1), (10, 0), (11, 901)):
            ws.cell(2, col).value = value
        wb.save(self.main)
        wb.close()
        row = next(row for row in self.preview()['parts'] if row['part_number'] == 'EC-20128A-TAB')
        self.assertIsNone(row['physical_qty'])
        self.assertIn('較舊', row['blocked_reason'])

    def test_other_deduction_remains_outside_count_absorption(self):
        from app.services.defective_deduction import deduct_defectives_from_main
        self.prepare_count_example()
        deduct_defectives_from_main(str(self.main), [{'part_number': 'EC-20128A-TAB', 'defective_qty': 3}],
                                   entry_header='其他調整扣帳')
        row = self.commit()['parts'][0]
        self.assertEqual(row['defect_delta'], -20)
        self.assertEqual(row['main_after'], 47)

    def test_post_count_new_defect_deducts_until_next_count(self):
        from app.services.defective_deduction import deduct_defectives_from_main
        from app.services.inventory_restore_guard import is_count_protected_record
        self.prepare_count_example()
        self.commit()
        deduct_defectives_from_main(str(self.main), [{'part_number': 'EC-20128A-TAB', 'defective_qty': 7}])
        new_id = db.create_defective_record(dict(part_number='EC-20128A-TAB', defective_qty=7))
        self.assertEqual(read_stock(str(self.main))['EC-20128A-TAB'], 43)
        self.assertFalse(is_count_protected_record(db.get_defective_record(new_id)))
        self.restart_session()
        result = self.commit()['parts'][0]
        self.assertEqual(result['main_after'], 50)
        self.assertEqual(result['main_adjustment'], 7)

    def test_old_replay_returns_400_without_touching_main_or_records(self):
        from app.routers import defectives
        from app.models import DefectiveReplayRequest
        self.prepare_count_example()
        # 真實 replay 查詢會 join 批次，讓測試紀錄完整。
        batch_id = db.create_defective_batch('existing-defect.xlsx')
        self.conn.execute('UPDATE defective_records SET batch_id=?', (batch_id,))
        self.commit()
        before = self.main.read_bytes()
        records = db.get_defective_records()
        with patch.object(defectives, 'backup_main_file') as backup:
            with self.assertRaises(HTTPException) as caught:
                asyncio.run(defectives.replay_after_rollback(DefectiveReplayRequest(cutoff='2026-09-19T10:00:00')))
        self.assertEqual(caught.exception.status_code, 400)
        backup.assert_not_called()
        self.assertEqual(self.main.read_bytes(), before)
        self.assertEqual(db.get_defective_records(), records)

    def test_parser_prefers_fg_and_preserves_exact_tab(self):
        parsed = parse_st_reconcile_file(str(self.count))
        self.assertEqual(parsed['source_columns'], {'book': 'F', 'physical': 'G'})
        self.assertEqual(parsed['rows'][0]['part_number'], 'EC-20128A-TAB')
        self.assertEqual(parsed['rows'][0]['physical_qty'], 771)
        self.assertEqual(parsed['count_date'], '2026-09-23')

    def restart_session(self):
        session = db.get_active_inventory_count_session('st')
        if session:
            db.finish_inventory_count_session(session['id'], status='cancelled')
        db.start_inventory_count_session(cutoff_at='2026-09-19T10:00:00', cutoff_code='9-12')

    def test_repeat_count_uses_main_adjustments_and_only_corrects_changed_quantity(self):
        part = 'EC-30037A-TAB'
        self.assertEqual(self.commit(parts=[part])['parts'][0]['main_after'], 3968)
        self.restart_session()
        self.assertEqual(self.commit(parts=[part])['parts'][0]['main_adjustment'], 0)
        self.restart_session()
        wb = load_workbook(self.count)
        wb.active['G6'] = 361
        wb.save(self.count)
        wb.close()
        result = self.commit(parts=[part])['parts'][0]
        self.assertEqual(result['main_adjustment'], 10)
        self.assertEqual(result['main_after'], 3978)
        self.assertEqual(read_stock(str(self.main))['EC-20128A-TAB'], 635)

    def test_partial_commit_allows_other_parts_and_older_count_only_blocks_counted_part(self):
        self.commit(parts=['EC-30037A-TAB'])
        self.restart_session()
        self.assertEqual(self.commit(parts=['EC-20128A-TAB'])['parts'][0]['main_after'], 626)
        self.restart_session()
        wb = load_workbook(self.count)
        wb.active['A1'] = '盤點日期 2026/9/22'
        wb.save(self.count)
        wb.close()
        rows = {row['part_number']: row for row in self.preview()['parts']}
        self.assertIsNone(rows['EC-30037A-TAB']['physical_qty'])
        self.assertIn('較舊', rows['EC-30037A-TAB']['blocked_reason'])
        self.assertEqual(rows['PART-KEEP']['physical_qty'], 50)
        before = self.file_hash(self.main)
        with self.assertRaisesRegex(ValueError, '未填實盤'):
            self.commit(parts=['EC-30037A-TAB'])
        self.assertEqual(self.file_hash(self.main), before)

    def test_restored_main_does_not_retain_adjustment_from_later_file(self):
        original = self.main.read_bytes()
        self.commit(parts=['EC-30037A-TAB'])
        self.main.write_bytes(original)
        self.restart_session()
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-30037A-TAB')
        self.assertEqual(row['prior_adjustment'], 0)
        self.assertEqual(row['main_adjustment'], 3199)

    def test_alias_sources_preserve_blank_rows_and_never_commit_partial_sum(self):
        for duplicate_part in ('EC-20128A', 'EC-20128A-TAB'):
            with self.subTest(source=duplicate_part):
                wb = load_workbook(self.count)
                wb.active.cell(8, 4, duplicate_part)
                wb.active.cell(8, 6, 10)
                wb.save(self.count)
                wb.close()
                row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
                self.assertEqual(len(row['source_rows']), 2)
                self.assertEqual([r['physical_qty'] for r in row['source_rows']], [771, None])
                self.assertIsNone(row['physical_qty'])
                self.assertTrue(row['warnings'])
                before = self.file_hash(self.main)
                with self.assertRaisesRegex(ValueError, '未填實盤'):
                    self.commit()
                self.assertEqual(self.file_hash(self.main), before)

    def test_complete_alias_sources_warn_and_show_each_quantity(self):
        wb = load_workbook(self.count)
        wb.active.cell(8, 4, 'EC-20128A')
        wb.active.cell(8, 6, 10)
        wb.active.cell(8, 7, 10)
        wb.save(self.count)
        wb.close()
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['physical_qty'], 781)
        self.assertEqual(len(row['source_rows']), 2)
        self.assertTrue(row['warnings'])

    def test_main_change_during_preview_is_rejected(self):
        original = st_reconcile.read_batch_stock
        def changed_main(*args):
            result = original(*args)
            wb = load_workbook(self.main)
            wb.active['T2'] = 123
            wb.save(self.main)
            wb.close()
            return result
        with patch.object(st_reconcile, 'read_batch_stock', side_effect=changed_main):
            with self.assertRaisesRegex(ValueError, '試算期間主檔已變更'):
                self.preview()

    def test_legacy_reversal_uses_actual_main_amount_without_guessing_db_history(self):
        wb = load_workbook(self.main)
        ws = wb.active
        for col, value in ((21, '不良品回復'), (22, '使用數量'), (23, '結存')):
            ws.cell(1, col, value)
        for col, value in ((21, 5), (22, 0), (23, 640)):
            ws.cell(2, col, value)
        wb.save(self.main)
        wb.close()
        rows = {row['part_number']: row for row in self.preview()['parts']}
        self.assertEqual(rows['EC-20128A-TAB']['defect_delta'], 5)
        self.assertEqual(rows['EC-20128A-TAB']['target_main'], 626)
        self.assertEqual(rows['EC-30037A-TAB']['main_adjustment'], 3199)

    def prepare_defective_reversal(self):
        self.restart_session()
        db.finish_inventory_count_session(db.get_active_inventory_count_session('st')['id'], status='cancelled')
        self.conn.execute('DELETE FROM defective_records')
        wb = load_workbook(self.main)
        for col, value in ((18, '不良品扣帳'), (19, '使用數量'), (20, '結存')):
            wb.active.cell(1, col).value = value
        for col, value in ((18, 0), (19, 5), (20, 775)):
            wb.active.cell(2, col).value = value
        wb.save(self.main)
        wb.close()
        wb = load_workbook(self.count)
        wb.active['G5'] = 775
        wb.save(self.count)
        wb.close()
        with patch.object(db, '_now', return_value='2026-09-20T10:00:00'):
            batch_id = db.create_defective_batch('defect.xlsx', main_file_mtime=self.main.stat().st_mtime)
            db.create_defective_record(dict(batch_id=batch_id, part_number='EC-20128A-TAB', defective_qty=5))
        return batch_id

    def delete_defective_batch(self, batch_id):
        from app.routers import defectives
        with patch.object(defectives, 'BACKUP_DIR', Path(self.tmp.name) / 'backups'), \
                patch.object(defectives, '_refresh_active_merge_drafts_after_main_change'):
            return asyncio.run(defectives.delete_batch(batch_id))

    def test_reversal_after_count_preserves_original_deduction_history(self):
        batch_id = self.prepare_defective_reversal()
        with patch.object(db, '_now', return_value='2026-09-24T10:00:00'):
            result = self.delete_defective_batch(batch_id)
        self.assertEqual(result['reversed_count'], 1)
        self.assertEqual(db.get_defective_batches(), [])
        self.restart_session()
        row = self.commit()['parts'][0]
        self.assertEqual(row['expected_count'], 780)
        self.assertEqual(row['main_after'], 775)
        self.assertEqual(row['main_adjustment'], -5)

    def test_reversal_before_count_and_restore_use_present_main_events_only(self):
        batch_id = self.prepare_defective_reversal()
        before_reversal = self.main.read_bytes()
        with patch.object(db, '_now', return_value='2026-09-22T10:00:00'):
            self.delete_defective_batch(batch_id)
        self.restart_session()
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['defect_delta'], 0)
        self.main.write_bytes(before_reversal)
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['defect_delta'], -5)
        self.assertEqual(row['main_adjustment'], 0)
        self.conn.execute('UPDATE defective_deleted_history SET main_period_id=999')
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['defect_delta'], -5)

    def test_reversal_archive_failure_rolls_back_main_snapshot_and_keeps_records(self):
        batch_id = self.prepare_defective_reversal()
        before = self.file_hash(self.main)
        snapshot = db.capture_inventory_snapshot_state()
        self.conn.execute("CREATE TRIGGER reject_history BEFORE INSERT ON defective_deleted_history BEGIN SELECT RAISE(ABORT, 'archive failed'); END")
        with self.assertRaisesRegex(sqlite3.IntegrityError, 'archive failed'):
            self.delete_defective_batch(batch_id)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertEqual(db.capture_inventory_snapshot_state(), snapshot)
        self.assertEqual(len(db.get_defective_batches()), 1)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM defective_deleted_history').fetchone()[0], 0)

    def test_reversal_save_failure_keeps_original_main_and_database(self):
        from app.services import defective_deduction
        batch_id = self.prepare_defective_reversal()
        before = self.file_hash(self.main)
        with patch.object(defective_deduction, '_save_workbook_atomically', side_effect=OSError('save failed')):
            with self.assertRaisesRegex(OSError, 'save failed'):
                self.delete_defective_batch(batch_id)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertEqual(len(db.get_defective_batches()), 1)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM defective_deleted_history').fetchone()[0], 0)

    def test_same_part_records_reverse_total_and_post_commit_hook_failure_still_succeeds(self):
        from app.routers import defectives
        batch_id = self.prepare_defective_reversal()
        with patch.object(db, '_now', return_value='2026-09-21T10:00:00'):
            db.create_defective_record(dict(batch_id=batch_id, part_number='EC-20128A-TAB', defective_qty=4))
        wb = load_workbook(self.main)
        wb.active['S2'] = 9
        wb.active['T2'] = 771
        wb.save(self.main)
        wb.close()
        wb = load_workbook(self.count)
        wb.active['G5'] = 771
        wb.save(self.count)
        wb.close()
        self.conn.execute('UPDATE defective_batches SET main_file_mtime=0 WHERE id=?', (batch_id,))
        with patch.object(defectives, 'BACKUP_DIR', Path(self.tmp.name) / 'backups'), \
                patch.object(db, '_now', return_value='2026-09-24T10:00:00'), \
                patch.object(defectives, '_refresh_active_merge_drafts_after_main_change', side_effect=RuntimeError('draft failed')):
            result = asyncio.run(defectives.delete_batch(batch_id))
        self.assertTrue(result['ok'])
        self.assertEqual(read_stock(str(self.main))['EC-20128A-TAB'], 780)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM defective_deleted_history').fetchone()[0], 2)
        self.restart_session()
        row = self.commit()['parts'][0]
        self.assertEqual(row['main_after'], 771)
        self.assertEqual(row['defect_delta'], 0)

    def test_db_only_deduction_does_not_affect_main_reconciliation(self):
        record_id = self.conn.execute("SELECT id FROM defective_records WHERE defective_qty=4").fetchone()[0]
        before = self.file_hash(self.main)
        db.delete_defective_record(record_id)
        self.assertIsNone(db.get_defective_record(record_id))
        row = next(r for r in self.preview()['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['defect_delta'], 0)
        self.assertEqual(row['main_adjustment'], -9)
        self.assertEqual(self.file_hash(self.main), before)

    def test_preview_uses_main_ledger_and_st_is_warning_only(self):
        rows = {row['part_number']: row for row in self.preview()['parts']}
        ec = rows['EC-20128A-TAB']
        self.assertEqual(
            [ec[key] for key in ('cutoff_main', 'defect_delta', 'expected_count', 'current_main', 'preserved_delta', 'target_main', 'main_adjustment')],
            [780, 0, 780, 635, -145, 626, -9],
        )
        self.assertEqual(ec['current_st'], 22000)
        other = rows['EC-30037A-TAB']
        self.assertEqual(
            [other[key] for key in ('cutoff_main', 'expected_count', 'physical_qty', 'current_main', 'target_main', 'main_adjustment')],
            [-2848, -2848, 351, 769, 3968, 3199],
        )
        self.assertEqual(other['current_st'], 9236)

    def test_missing_part_suggestions_use_valid_unoccupied_main_balance_not_st(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        self.conn.execute("INSERT INTO st_inventory_snapshot(part_number,stock_qty) VALUES('EC-20128A-TA8',12345)")
        report = self.preview()
        missing = report['uncovered_parts'][0]
        self.assertEqual(missing['source_part_number'], 'EC-20128A-TA8')
        self.assertTrue(missing['can_map'])
        self.assertEqual(missing['suggestions'], [{'part_number': 'EC-20128A-TAB', 'stock_qty': 635}])
        self.assertEqual(report['part_mappings'], {})

    def test_st_only_exact_part_does_not_block_main_tab_alias(self):
        self.set_count_part(5, 'EC-20128A')
        self.conn.execute("INSERT INTO st_inventory_snapshot(part_number,stock_qty) VALUES('EC-20128A',12345)")
        row = next(row for row in self.preview()['parts'] if row['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['source_part_numbers'], ['EC-20128A'])
        self.assertEqual(row['expected_count'], 780)
        self.assertNotIn('manual_mapping', row)

    def test_mapping_normalizes_and_uses_target_defects_and_main_balances(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        report = self.preview({' ec-20128a-ta8 ': ' ec-20128a-tab '})
        self.assertEqual(report['part_mappings'], {'EC-20128A-TA8': 'EC-20128A-TAB'})
        row = next(row for row in report['parts'] if row['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['source_part_numbers'], ['EC-20128A-TA8'])
        self.assertTrue(row['manual_mapping'])
        self.assertIn('EC-20128A-TA8 → EC-20128A-TAB', row['warnings'][0])
        self.assertEqual([row[key] for key in ('physical_qty', 'defect_delta', 'expected_count', 'target_main', 'main_adjustment')], [771, 0, 780, 626, -9])
        self.assertEqual(report['uncovered_parts'], [])

    def test_mapped_commit_only_writes_selected_target_and_preserves_source_stock(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        mappings = {'EC-20128A-TA8': 'EC-20128A-TAB'}
        before_st = db.get_st_inventory_stock()
        result = self.commit(mappings=mappings)
        self.assertEqual(result['part_mappings'], mappings)
        self.assertEqual(result['parts'][0]['source_part_numbers'], ['EC-20128A-TA8'])
        self.assertEqual(result['parts'][0]['main_after'], 626)
        self.assertEqual(db.get_st_inventory_stock(), before_st)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM st_reconcile_alignments').fetchone()[0], 0)
        self.assertEqual(read_stock(str(self.main))['EC-30037A-TAB'], 769)
        self.assertEqual(parse_st_reconcile_file(str(self.count))['rows'][0]['part_number'], 'EC-20128A-TA8')

    def test_exact_and_alias_matches_cannot_be_manually_overridden(self):
        with self.assertRaisesRegex(ValueError, '已匹配主檔'):
            self.preview({'EC-20128A-TAB': 'PART-KEEP'})
        self.set_count_part(5, 'EC-20128A')
        with self.assertRaisesRegex(ValueError, '已匹配主檔'):
            self.preview({'EC-20128A': 'PART-KEEP'})

    def test_existing_main_part_without_valid_balance_cannot_be_mapped(self):
        self.invalidate_cutoff(2)
        row = self.preview()['uncovered_parts'][0]
        self.assertFalse(row['can_map'])
        self.assertEqual(row['suggestions'], [])
        self.assertIn('請先檢查主檔', row['reason'])
        with self.assertRaisesRegex(ValueError, '已匹配主檔'):
            self.preview({'EC-20128A-TAB': 'PART-KEEP'})

    def test_target_requires_effective_main_cutoff_and_current_balance(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        self.invalidate_cutoff(2)
        with self.assertRaisesRegex(ValueError, '有效截止批次及目前結存'):
            self.preview({'EC-20128A-TA8': 'EC-20128A-TAB'})
        self.assertNotIn('EC-20128A-TAB', [item['part_number'] for item in self.preview()['uncovered_parts'][0]['suggestions']])
        with self.assertRaisesRegex(ValueError, '有效截止批次及目前結存'):
            self.preview({'EC-20128A-TA8': 'ST-ONLY'})

    def test_mapping_rejects_unknown_sources_and_malformed_schema(self):
        invalid_mappings = [[], {'EC-20128A-TAB': 1}, {'': 'PART-KEEP'}, {'EC-20128A-TAB': ''}, {' ec-20128a-tab ': 'PART-KEEP', 'EC-20128A-TAB': 'PART-KEEP'}]
        for mappings in invalid_mappings:
            with self.subTest(mappings=mappings), self.assertRaises(ValueError):
                self.preview(mappings)
        with self.assertRaisesRegex(ValueError, '不在本次盤點表'):
            self.preview({'NOT-IN-FILE': 'EC-20128A-TAB'})

    def test_mapping_cannot_merge_different_sources_or_occupied_exact_target(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        with self.assertRaisesRegex(ValueError, '不會自動合併'):
            self.preview({'EC-20128A-TA8': 'PART-KEEP'})
        self.set_count_part(6, 'EC-20128A-TA9')
        with self.assertRaisesRegex(ValueError, '不會自動合併'):
            self.preview({'EC-20128A-TA8': 'EC-20128A-TAB', 'EC-20128A-TA9': 'EC-20128A-TAB'})
        self.set_count_part(6, 'EC-20128A')
        with self.assertRaisesRegex(ValueError, '不會自動合併'):
            self.preview({'EC-20128A-TA8': 'EC-20128A-TAB'})

    def test_same_source_duplicate_rows_keep_original_sum_after_mapping(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        self.set_count_part(6, 'EC-20128A-TA8')
        row = next(row for row in self.preview({'EC-20128A-TA8': 'EC-20128A-TAB'})['parts'] if row['part_number'] == 'EC-20128A-TAB')
        self.assertEqual(row['source_part_numbers'], ['EC-20128A-TA8'])
        self.assertEqual(row['physical_qty'], 1122)
        self.assertEqual(row['book_qty'], -2068)

    def test_mapping_change_requires_new_preview_token_and_writes_nothing(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        token = self.preview()['preview_token']
        before = self.file_hash(self.main)
        mappings = {'EC-20128A-TA8': 'EC-20128A-TAB'}
        self.assertNotEqual(token, self.preview(mappings)['preview_token'])
        with self.assertRaisesRegex(StaleReconcilePreviewError, '重新試算'):
            self.commit(token=token, mappings=mappings)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))
        self.assertFalse((Path(self.tmp.name) / 'backups').exists())

    def test_invalid_mapping_commit_is_input_error_and_does_not_write(self):
        before = self.file_hash(self.main)
        with self.assertRaises(st_reconcile.PartMappingError):
            self.commit(mappings={'EC-20128A-TAB': 'PART-KEEP'}, token=self.preview()['preview_token'])
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))

    def test_legacy_nonempty_mappings_are_rejected_without_writes(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        before_st = db.get_st_inventory_stock()
        mappings = {'EC-20128A-TA8': 'EC-20128A-TAB'}
        with self.assertRaisesRegex(ValueError, '只支援選擇截止批次'):
            build_st_reconcile_preview(str(self.count), '2026-09-19', part_mappings=mappings)
        with self.assertRaisesRegex(ValueError, '只支援選擇截止批次'):
            commit_st_reconcile_stop_loss(str(self.count), '2026-09-19', part_mappings=mappings)
        self.assertEqual(db.get_st_inventory_stock(), before_st)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM st_reconcile_alignments').fetchone()[0], 0)

    def test_api_json_mapping_validation_precedes_file_reads(self):
        invalid_json = ['not json', '[]', 'null', '"text"', '{"A": 1}', '{"A": "B", "A": "C"}', '{" a ": "B", "A": "C"}']
        for raw in invalid_json:
            for endpoint in (reconcile_router.preview_st_reconcile, reconcile_router.commit_st_reconcile):
                upload = UploadFile(filename=self.count.name, file=io.BytesIO(self.count.read_bytes()))
                with self.subTest(raw=raw, endpoint=endpoint.__name__), self.assertRaises(HTTPException) as raised:
                    asyncio.run(endpoint(file=upload, part_mappings=raw))
                self.assertEqual(raised.exception.status_code, 400)
                self.assertEqual(upload.file.tell(), 0)

    def test_api_form_mapping_is_forwarded_and_token_change_returns_409(self):
        self.set_count_part(5, 'EC-20128A-TA8')
        app = FastAPI()
        app.include_router(reconcile_router.router)
        old_token = self.preview()['preview_token']
        before = self.file_hash(self.main)
        with (
            TestClient(app) as client,
            patch.object(reconcile_router, '_resolve_cutoff', return_value=('2026-09-19T10:00:00', '9-12')),
            patch.object(reconcile_router, 'RECONCILE_UPLOAD_DIR', Path(self.tmp.name)),
        ):
            form = {'cutoff_batch_code': '9-12', 'part_mappings': json.dumps({' ec-20128a-ta8 ': ' ec-20128a-tab '})}
            files = {'file': (self.count.name, self.count.read_bytes(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            response = client.post('/reconcile/st/preview', data=form, files=files)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['part_mappings'], {'EC-20128A-TA8': 'EC-20128A-TAB'})
            form.update(part_numbers='EC-20128A-TAB', preview_token=old_token)
            response = client.post('/reconcile/st/commit', data=form, files=files)
            self.assertEqual(response.status_code, 409, response.text)
            self.assertIn('重新試算', response.json()['detail'])
            form['part_mappings'] = json.dumps({'EC-20128A-TAB': 'PART-KEEP'})
            response = client.post('/reconcile/st/commit', data=form, files=files)
            self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))
        self.assertFalse((Path(self.tmp.name) / 'backups').exists())

    def test_commit_writes_only_selected_main_and_never_st_or_alignment(self):
        st_before = db.get_st_inventory_stock()
        wb = load_workbook(self.main, data_only=False)
        untouched_before = [cell.value for cell in wb.active[4]]
        wb.close()

        result = self.commit(parts=['EC-30037A-TAB'])

        self.assertTrue(Path(result['backup_path']).is_file())
        self.assertEqual(result['parts'][0]['main_before'], 769)
        self.assertEqual(result['parts'][0]['main_after'], 3968)
        self.assertEqual(result['parts'][0]['main_adjustment'], 3199)
        self.assertEqual(db.get_st_inventory_stock(), st_before)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM st_reconcile_alignments').fetchone()[0], 0)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM st_inventory_audit_log').fetchone()[0], 0)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM defective_records WHERE absorbed_by_alignment_id<>0').fetchone()[0], 0)
        self.assertIsNone(db.get_active_inventory_count_session('st'))

        wb = load_workbook(self.main, data_only=False)
        ws = wb.active
        self.assertIn('盤點調整', ws.cell(1, 18).value)
        self.assertEqual([ws.cell(3, col).value for col in (18, 19, 20)], [3199, 0, 351])
        untouched_after = [cell.value for cell in ws[4]]
        self.assertEqual(untouched_after[:17] + untouched_after[20:], untouched_before)
        self.assertEqual(untouched_after[17:20], [None, None, None])
        wb.close()
        self.assertEqual(read_stock(str(self.main))['EC-30037A-TAB'], 3968)
        self.assertEqual(db.get_snapshot_stock()['EC-30037A-TAB'], 3968)

    def test_zero_adjustment_inserts_event_and_recalculates_later_balance(self):
        wb = load_workbook(self.count)
        wb.active['G5'] = 780
        wb.save(self.count)
        wb.close()
        result = self.commit(parts=['EC-20128A-TAB'])
        self.assertEqual(result['parts'][0]['main_adjustment'], 0)
        self.assertEqual(result['parts'][0]['main_after'], 635)
        wb = load_workbook(self.main, data_only=False)
        ws = wb.active
        self.assertEqual([ws.cell(2, col).value for col in (18, 19, 20, 21, 22, 23)], [0, 0, 780, 0, 145, 635])
        wb.close()

    def test_nonzero_adjustment_recalculates_all_later_balances(self):
        wb = load_workbook(self.count)
        wb.active['G5'] = 770
        wb.save(self.count)
        wb.close()
        result = self.commit(parts=['EC-20128A-TAB'])
        self.assertEqual(result['parts'][0]['main_adjustment'], -10)
        self.assertEqual(result['parts'][0]['main_after'], 625)
        wb = load_workbook(self.main, data_only=False)
        self.assertEqual([wb.active.cell(2, col).value for col in (18, 19, 20, 23)], [0, 10, 770, 625])
        wb.close()

    def test_stale_token_writes_nothing_and_session_stays_active(self):
        token = self.preview()['preview_token']
        wb = load_workbook(self.main)
        wb.active['T2'] = 630
        wb.save(self.main)
        wb.close()
        before = self.file_hash(self.main)
        with self.assertRaisesRegex(StaleReconcilePreviewError, '重新試算'):
            self.commit(parts=['EC-20128A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))

    def test_formula_in_later_input_is_rejected_before_backup_or_write(self):
        wb = load_workbook(self.main)
        wb.active['R2'] = '=1+1'
        wb.save(self.main)
        wb.close()
        before = self.file_hash(self.main)
        token = self.preview()['preview_token']
        with self.assertRaisesRegex(ValueError, '含公式'):
            self.commit(parts=['EC-20128A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertFalse((Path(self.tmp.name) / 'backups').exists())

    def test_formula_in_later_balance_is_preserved_and_writes_nothing(self):
        wb = load_workbook(self.main)
        wb.active['T2'] = '=Q2+R2-S2'
        wb.save(self.main)
        wb.close()
        before = self.file_hash(self.main)
        token = self.preview()['preview_token']
        with self.assertRaisesRegex(ValueError, '含公式'):
            self.commit(parts=['EC-20128A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        saved = load_workbook(self.main, data_only=False)
        self.assertEqual(saved.active['T2'].value, '=Q2+R2-S2')
        saved.close()
        self.assertFalse((Path(self.tmp.name) / 'backups').exists())

    def test_snapshot_failure_restores_main_and_keeps_session_active(self):
        before = self.file_hash(self.main)
        token = self.preview()['preview_token']
        snapshot_before = db.capture_inventory_snapshot_state()
        with patch.object(st_reconcile, 'refresh_snapshot_from_main', side_effect=RuntimeError('snapshot failed')):
            with self.assertRaisesRegex(RuntimeError, 'snapshot failed'):
                self.commit(parts=['EC-30037A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))
        self.assertEqual(db.capture_inventory_snapshot_state(), snapshot_before)

    def test_finish_session_failure_restores_main_and_snapshot(self):
        before = self.file_hash(self.main)
        token = self.preview()['preview_token']
        snapshot_before = db.capture_inventory_snapshot_state()
        with patch.object(db, 'finish_inventory_count_session', return_value=False):
            with self.assertRaisesRegex(RuntimeError, '工作階段完成失敗'):
                self.commit(parts=['EC-30037A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))
        self.assertEqual(db.capture_inventory_snapshot_state(), snapshot_before)

    def test_preview_rebuild_failure_is_stale_and_writes_nothing(self):
        token = self.preview()['preview_token']
        wb = load_workbook(self.main)
        ws = wb.active
        ws['L1'] = 'changed'
        ws['O1'] = 'changed'
        wb.save(self.main)
        wb.close()
        before = self.file_hash(self.main)
        with self.assertRaisesRegex(StaleReconcilePreviewError, '重新試算'):
            self.commit(parts=['EC-20128A-TAB'], token=token)
        self.assertEqual(self.file_hash(self.main), before)
        self.assertIsNotNone(db.get_active_inventory_count_session('st'))

    def test_missing_part_anchor_or_selection_writes_nothing(self):
        token = self.preview()['preview_token']
        before = self.file_hash(self.main)
        with self.assertRaisesRegex(ValueError, '至少勾選'):
            commit_st_reconcile_stop_loss(
                str(self.count), '2026-09-19T10:00:00', cutoff_label='9-12',
                part_numbers=[], preview_token=token,
            )
        with self.assertRaisesRegex(ValueError, '不可對帳'):
            self.commit(parts=['MISSING'], token=token)
        self.assertEqual(self.file_hash(self.main), before)

    def test_api_batch_commit_rejects_empty_part_numbers_before_read(self):
        upload = UploadFile(filename=self.count.name, file=io.BytesIO(self.count.read_bytes()))
        with patch.object(reconcile_router, '_resolve_cutoff', return_value=('2026-09-19T10:00:00', '9-12')):
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(reconcile_router.commit_st_reconcile(
                    cutoff_batch_code='9-12',
                    file=upload,
                    part_numbers=None,
                    preview_token='token',
                ))
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn('至少勾選', raised.exception.detail)

    def test_api_activity_log_failure_does_not_turn_success_into_failure(self):
        upload = UploadFile(filename=self.count.name, file=io.BytesIO(self.count.read_bytes()))
        expected = {
            'ok': True,
            'session_completed': True,
            'summary': {'session_id': 1, 'part_count': 1},
        }
        with (
            patch.object(reconcile_router, '_resolve_cutoff', return_value=('2026-09-19T10:00:00', '9-12')),
            patch.object(reconcile_router, 'commit_st_reconcile_stop_loss', return_value=expected),
            patch.object(db, 'log_activity', side_effect=RuntimeError('log failed')),
            patch.object(reconcile_router.log, 'exception'),
        ):
            result = asyncio.run(reconcile_router.commit_st_reconcile(
                cutoff_batch_code='9-12',
                file=upload,
                part_numbers=['EC-20128A-TAB'],
                preview_token='token',
            ))
        self.assertEqual(result, expected)

    def test_api_stale_preview_returns_409(self):
        upload = UploadFile(filename=self.count.name, file=io.BytesIO(self.count.read_bytes()))
        with (
            patch.object(reconcile_router, '_resolve_cutoff', return_value=('2026-09-19T10:00:00', '9-12')),
            patch.object(
                reconcile_router,
                'commit_st_reconcile_stop_loss',
                side_effect=StaleReconcilePreviewError('主檔已變更，請重新試算'),
            ),
        ):
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(reconcile_router.commit_st_reconcile(
                    cutoff_batch_code='9-12',
                    file=upload,
                    part_numbers=['EC-20128A-TAB'],
                    preview_token='token',
                ))
        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn('重新試算', raised.exception.detail)

    def test_api_legacy_date_commit_does_not_require_preview_token(self):
        upload = UploadFile(filename=self.count.name, file=io.BytesIO(self.count.read_bytes()))
        expected = {
            'ok': True,
            'session_completed': True,
            'summary': {'session_id': 1, 'part_count': 1},
        }
        with (
            patch.object(reconcile_router, '_resolve_cutoff', return_value=('2026-09-19T10:00:00', '')),
            patch.object(reconcile_router, 'commit_st_reconcile_stop_loss', return_value=expected) as commit_mock,
            patch.object(db, 'log_activity'),
        ):
            result = asyncio.run(reconcile_router.commit_st_reconcile(
                cutoff_date='2026-09-19',
                file=upload,
                part_numbers=['EC-20128A-TAB'],
                preview_token='',
            ))
        self.assertEqual(result, expected)
        self.assertEqual(commit_mock.call_args.kwargs['preview_token'], '')

    def test_blank_g_does_not_fall_back_to_total_and_filename_date_is_detected(self):
        wb = load_workbook(self.count)
        ws = wb.active
        ws['A1'] = None
        for row in range(5, 8):
            ws.cell(row, 6).value = None
            ws.cell(row, 7).value = None
        wb.save(self.count)
        wb.close()
        parsed = parse_st_reconcile_file(str(self.count))
        self.assertEqual(parsed['source_columns'], {'book': 'V', 'physical': 'G'})
        self.assertIsNone(parsed['rows'][0]['physical_qty'])
        self.assertEqual(parsed['count_date'], '2026-09-23')

    def test_same_batch_rightmost_valid_and_blank_falls_back(self):
        wb = load_workbook(self.main)
        ws = wb.active
        ws['Q1'] = 'MODEL-X'
        ws['Q2'] = None
        wb.save(self.main)
        stock = read_batch_stock(str(self.main), '9-12')['EC-20128A-TAB']
        self.assertEqual(stock, {'cutoff_main': 800, 'current_main': 635})
        ws['N2'] = None
        ws['P2'] = 99999
        wb.save(self.main)
        wb.close()
        stock = read_batch_stock(str(self.main), '9-12')['EC-20128A-TAB']
        self.assertEqual(stock['cutoff_main'], 890)


if __name__ == '__main__':
    unittest.main()
