from __future__ import annotations

import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from app.services.bom_sorting import sort_bom_sections
from app.services.merge_drafts import _write_draft_files
from main import app


class BomSortingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("PYTEST_CURRENT_TEST", "1")

    def _build_workbook(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "BOM"
        ws["G2"] = 10
        ws["A4"] = "自備"
        ws["A8"] = "CONSIGN-PRODUCTION"
        ws["A11"] = "OPENTEXT PURCHASING PARTS"
        ws["A14"] = "耗材紙箱"
        for row, part, quantity in [
            (5, "Z-2", 2), (6, "A-1", 3), (7, "A-1", 4),
            (9, "Z-2", 5), (10, "A-0", 6),
            (12, "P-2", 7), (13, "P-1", 8),
            (15, "K-2", 9), (16, "K-1", 10),
        ]:
            ws.cell(row, 1, row)
            ws.cell(row, 2, quantity)
            ws.cell(row, 3, part)
            ws.cell(row, 4, f"明細-{row}")
            ws.cell(row, 6, f"=B{row}*G$2")
            ws.cell(row, 7, row + 20)
            ws.cell(row, 8, row + 30)
            ws.cell(row, 9, f"=SUM(G{row},H{row})")
            ws.cell(row, 10, f"=I{row}-F{row}")
            ws.cell(row, 11, f"備註-{row}")
            ws.row_dimensions[row].height = row + 15
        ws["C5"].comment = Comment("Z-2 備註", "Andy")
        ws["C5"].hyperlink = "https://example.com/Z-2"
        ws["H5"].fill = PatternFill(fill_type="solid", fgColor="FFFFC000")
        ws["B5"].number_format = "0.00"
        ws.merge_cells("K5:L5")
        ws.merge_cells("K10:L10")
        ws["I8"] = "=SUM(G8,H8)"
        ws["J8"] = "=I8-F8"
        return workbook

    def _assert_sorted_sections(self, ws):
        self.assertEqual([ws.cell(row, 3).value for row in (5, 6, 7)], ["A-1", "A-1", "Z-2"])
        self.assertEqual([ws.cell(row, 3).value for row in (9, 10)], ["A-0", "Z-2"])
        self.assertEqual([ws.cell(row, 3).value for row in (12, 13)], ["P-1", "P-2"])
        self.assertEqual([ws.cell(row, 3).value for row in (15, 16)], ["K-1", "K-2"])
        self.assertEqual([ws.cell(row, 4).value for row in (5, 6, 7, 9, 10)], ["明細-6", "明細-7", "明細-5", "明細-10", "明細-9"])
        self.assertEqual(ws["A8"].value, "CONSIGN-PRODUCTION")
        self.assertEqual(ws["A11"].value, "OPENTEXT PURCHASING PARTS")
        self.assertEqual(ws["A14"].value, "耗材紙箱")
        self.assertEqual(ws["I8"].value, "=SUM(G8,H8)")
        self.assertEqual(ws["F7"].value, "=B7*G$2")
        self.assertEqual(ws["I7"].value, "=SUM(G7,H7)")
        self.assertEqual(ws["J7"].value, "=I7-F7")
        self.assertEqual(ws["B7"].value, 2)
        self.assertEqual(ws["B7"].number_format, "0.00")
        self.assertEqual(ws["K7"].value, "備註-5")
        self.assertEqual(ws["C7"].comment.text, "Z-2 備註")
        self.assertEqual(ws["C7"].hyperlink.ref, "C7")
        self.assertEqual(ws["H7"].fill.fgColor.rgb, "FFFFC000")
        self.assertEqual(ws.row_dimensions[7].height, 20)
        self.assertEqual({str(area) for area in ws.merged_cells.ranges}, {"K7:L7", "K9:L9"})

    def test_sections_preserve_rows_formulas_styles_and_duplicate_order_after_save(self):
        workbook = self._build_workbook()
        sort_bom_sections(workbook.active)
        sort_bom_sections(workbook.active)
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        result = load_workbook(buffer)
        try:
            self._assert_sorted_sections(result.active)
            self.assertEqual(result.active["G7"].value, 25)
            self.assertEqual(result.active["H7"].value, 35)
        finally:
            result.close()

    def test_blank_and_merged_headings_and_vertical_labels_keep_sections_separate(self):
        workbook = Workbook()
        ws = workbook.active
        for row, part in [(5, "Z"), (6, "B"), (8, "Y"), (9, "A"), (11, "X"), (12, "C"), (13, "W"), (14, "D")]:
            ws.cell(row, 3, part)
        ws.merge_cells("A8:A9")
        ws["A8"] = "客供"
        ws.merge_cells("C10:F10")
        ws["C10"] = "自備料"
        ws.merge_cells("A11:A12")
        ws["A11"] = "自備"
        sort_bom_sections(ws)
        self.assertEqual([ws.cell(row, 3).value for row in range(5, 15)], ["B", "Z", None, "A", "Y", "自備料", "C", "X", "D", "W"])
        self.assertEqual(ws["A8"].value, "客供")
        self.assertEqual(ws["A11"].value, "自備")
        workbook.close()

    def test_draft_generation_sorts_every_file_without_changing_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            workbook = self._build_workbook()
            workbook.create_sheet("說明")["A1"] = "原始說明"
            workbook.active = 1
            workbook.save(source)
            workbook.close()
            original_bytes = source.read_bytes()
            record = {"id": "bom-1", "filepath": str(source), "filename": source.name}
            with patch("app.services.merge_drafts.db.get_bom_file", return_value=record), \
                 patch("app.services.merge_drafts._ensure_editable_bom_for_draft", return_value=record), \
                 patch("app.services.merge_drafts.save_workbook_with_recalc", side_effect=lambda wb, path: wb.save(path)):
                written = _write_draft_files(1, [
                    {"bom_file_id": "bom-1", "model": "MODEL-A", "supplements": {"Z-2": 11}, "carry_overs": {"Z-2": 9}},
                    {"bom_file_id": "bom-1", "model": "MODEL-B", "supplements": {"Z-2": 12}},
                ], root_dir=root / "drafts")
            self.assertEqual(len(written), 2)
            for entry, supplement in zip(written, (11, 12)):
                result = load_workbook(entry["filepath"])
                try:
                    self._assert_sorted_sections(result.active)
                    self.assertEqual(result.active["H7"].value, supplement)
                    self.assertEqual(result.active["H10"].value, 0)
                    self.assertEqual(result["說明"]["A1"].value, "原始說明")
                finally:
                    result.close()
            self.assertEqual(source.read_bytes(), original_bytes)

    def test_direct_download_sorts_single_file_and_every_file_in_zip(self):
        client = TestClient(app)
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = self._build_workbook()
            workbook.save(source)
            workbook.close()
            original_bytes = source.read_bytes()
            records = [{"id": f"bom-{index}", "filepath": str(source), "filename": source.name, "model": f"MODEL-{index}"} for index in (1, 2)]
            for count in (1, 2):
                with self.subTest(files=count), \
                     patch("app.routers.bom.db.get_bom_files", return_value=records), \
                     patch("app.routers.bom._ensure_editable_bom_record", side_effect=lambda record: record), \
                     patch("app.routers.bom._build_direct_purchase_highlights", return_value=set()), \
                     patch("app.routers.bom._build_order_based_export_values", return_value=({}, {}, {}, {}, {})), \
                     patch("app.services.workbook_recalc.refresh_saved_workbook_formula_cache", return_value=False):
                    response = client.post("/api/bom/dispatch-download", json={"bom_ids": [record["id"] for record in records[:count]], "supplements": {"Z-2": 11}})
                    self.assertEqual(response.status_code, 200)
                    if count == 1:
                        files = [response.content]
                    else:
                        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                            files = [archive.read(name) for name in archive.namelist()]
                    self.assertEqual(len(files), count)
                    for content in files:
                        result = load_workbook(io.BytesIO(content))
                        try:
                            self._assert_sorted_sections(result.active)
                            self.assertEqual(result.active["H7"].value, 11)
                        finally:
                            result.close()
            self.assertEqual(source.read_bytes(), original_bytes)
