from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import test_st_reconcile_batch as batch_fixture
from app.services.main_reader import read_batch_stock, read_stock
from app.services.merge_to_main import merge_row_to_main, preview_order_batches, supplement_part_in_main


class InventoryFormulaFlowTests(unittest.TestCase):
    def setUp(self):
        self.case = batch_fixture.BatchReconcileTests()
        self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def edit(self, values):
        wb = load_workbook(self.case.main)
        for address, value in values.items():
            wb.active[address] = value
        wb.save(self.case.main)
        wb.close()

    def groups(self):
        return [{
            'batch_code': '9-14', 'bom_model': 'MODEL', 'po_number': 'PO',
            'components': [{'part_number': 'EC-20128A-TAB', 'needed_qty': 10, 'prev_qty_cs': 0}],
        }]

    def dispatch(self):
        return merge_row_to_main(str(self.case.main), self.groups(), {},
                                 backup_dir=str(Path(self.case.tmp.name) / 'dispatch_backups'))

    def test_formula_cutoff_preview_and_commit_preserve_matching_stock(self):
        self.edit({'Q2': '=N2+O2-P2'})
        preview = self.case.preview()
        row = next(r for r in preview['parts'] if r['part_number'] == 'EC-20128A-TAB')
        self.assertEqual((row['cutoff_main'], row['expected_count'], row['main_adjustment']), (780, 771, 0))
        self.case.commit(token=preview['preview_token'])
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 635)
        wb = load_workbook(self.case.main)
        self.assertEqual(wb.active['Q2'].value, '=N2+O2-P2')
        wb.close()

    def test_formula_current_supplement_then_dispatch_uses_current_value(self):
        self.edit({'T2': '=Q2+R2-S2'})
        result = supplement_part_in_main(str(self.case.main), 'EC-20128A-TAB', 25)
        self.assertEqual(result['stock_after'], 660)
        stocks = read_batch_stock(str(self.case.main), '9-12')['EC-20128A-TAB']
        self.assertEqual(stocks, {'cutoff_main': 780, 'current_main': 660})
        preview = preview_order_batches(str(self.case.main), [{'groups': self.groups()}])
        row = preview['batches'][0]['groups'][0]['rows'][0]
        self.assertEqual(row['j_value'], 650)
        self.dispatch()
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 650)
        wb = load_workbook(self.case.main)
        self.assertEqual(wb.active['T2'].value, '=Q2+R2-S2')
        wb.close()

    def test_invalid_cutoff_formula_blocks_preview_instead_of_using_older_balance(self):
        for value in ('=SUM(N2:P2)', '=Q2+1', '=1/0', '#REF!', 'NaN'):
            with self.subTest(value=value):
                self.edit({'Q2': value})
                before = self.case.main.read_bytes()
                with self.assertRaisesRegex(ValueError, r'EC-20128A-TAB.*Q2'):
                    self.case.preview()
                self.assertEqual(self.case.main.read_bytes(), before)

    def test_invalid_current_formula_blocks_dispatch_and_leaves_file_unchanged(self):
        self.edit({'T2': '=SUM(Q2:S2)'})
        before = self.case.main.read_bytes()
        with self.assertRaisesRegex(ValueError, r'EC-20128A-TAB.*T2'):
            self.dispatch()
        with self.assertRaisesRegex(ValueError, r'EC-20128A-TAB.*T2'):
            read_batch_stock(str(self.case.main), '9-12')
        self.assertEqual(self.case.main.read_bytes(), before)

    def test_blank_balance_skips_usage_and_unrelated_numeric_notes(self):
        self.edit({'T2': None, 'Z1': '備註', 'Z2': 999999})
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 780)
        self.assertEqual(read_batch_stock(str(self.case.main), '9-12')['EC-20128A-TAB']['current_main'], 780)
        self.dispatch()
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 770)

    def test_blank_batch_header_with_model_balance_uses_formula_consistently(self):
        self.edit({'R1': None, 'S1': 'PO#9', 'T1': 'MODEL-X', 'T2': '=Q2+R2-S2'})
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 635)
        self.assertEqual(read_batch_stock(str(self.case.main), '9-12')['EC-20128A-TAB']['current_main'], 635)
        self.dispatch()
        self.assertEqual(read_stock(str(self.case.main))['EC-20128A-TAB'], 625)

    def test_legacy_two_column_deductions_do_not_shift_later_balance_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['料號', '廠商', 'MOQ', None, None, None, None, '盤點',
                   None, '庚霖不良0127', None, 59220, 'TTUSBC-REV-A',
                   '不良品扣帳', '05/08 15:42', None,
                   '不良品扣帳', '05/25 10:38',
                   '9-12', 4500059403, 'MODEL'])
        ws.append(['EC-20128A-TAB', '', 4000, None, None, None, None, None,
                   None, None, 1123, 80, 1043, None, None, None, None, None,
                   None, None, None])
        ws.append(['PART-B', '', 1000, None, None, None, None, 0,
                   None, None, None, None, None, None, None, None, None, None,
                   1000, 300, 700])
        wb.save(self.case.main)
        wb.close()
        self.assertEqual(read_stock(str(self.case.main)), {'EC-20128A-TAB': 1043, 'PART-B': 700})
        self.assertEqual(read_batch_stock(str(self.case.main), '9-12'), {
            'EC-20128A-TAB': {'cutoff_main': 1043, 'current_main': 1043},
            'PART-B': {'cutoff_main': 700, 'current_main': 700},
        })
