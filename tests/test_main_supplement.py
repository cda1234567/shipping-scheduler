from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

from app.services.main_reader import read_stock
from app.services.merge_to_main import supplement_part_in_main


class MainSupplementTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "main.xlsx"

    def _build(self, values: list, *, legacy=False):
        wb = Workbook()
        ws = wb.active
        ws.append(["料號", "廠商", "MOQ", None, None, None, None, "期初庫存",
                   None if legacy else "10-1", "用量", "MODEL" if legacy else "結存",
                   "10-2", "用量", "結存"])
        ws.append(["PART-A", "", 100, None, None, None, None, *values])
        ws.append(["PART-B", "", 100, None, None, None, None,
                   50, 0, 1, "=H3+I3-J3", 0, 1, "=K3+L3-M3"])
        wb.save(self.path)
        wb.close()

    def _values(self):
        wb = load_workbook(self.path, data_only=False)
        try:
            return {cell.coordinate: cell.value for row in wb.active for cell in row}
        finally:
            wb.close()

    def _verify(self, qty: float, expected: float, allowed: set[str]):
        before_stock = read_stock(str(self.path))
        before_cells = self._values()
        result = supplement_part_in_main(str(self.path), "PART-A", qty)
        after_stock = read_stock(str(self.path))
        self.assertEqual(result["stock_before"], before_stock["PART-A"])
        self.assertEqual(result["stock_after"], expected)
        self.assertEqual(after_stock, {**before_stock, "PART-A": expected})
        after_cells = self._values()
        self.assertEqual(set(before_cells), set(after_cells))
        self.assertEqual({k: v for k, v in before_cells.items() if k not in allowed},
                         {k: v for k, v in after_cells.items() if k not in allowed})
        return result, after_cells

    def _cache_k2(self, value: float):
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(self.path) as source:
            entries = [(item, source.read(item.filename)) for item in source.infolist()]
        with zipfile.ZipFile(self.path, "w") as target:
            for item, data in entries:
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    cell = root.find(f".//{{{ns}}}c[@r='K2']")
                    cell.find(f"{{{ns}}}v").text = str(value)
                    data = ET.tostring(root)
                target.writestr(item, data)

    def test_uncached_balance_formula_updates_supply_and_keeps_usage_and_formula(self):
        self._build([100, 0, 125, "=H2+I2-J2"])
        result, cells = self._verify(25, 0, {"I2"})
        self.assertEqual(result["supplement_col"], 9)
        self.assertEqual(cells["I2"], 25)

    def test_formula_chain_updates_the_latest_batch_supply_only(self):
        self._build([100, 0, 125, "=H2+I2-J2", 0, 25, "=K2+L2-M2"])
        result, cells = self._verify(50, 0, {"L2"})
        self.assertEqual(result["supplement_col"], 12)
        self.assertEqual(cells["L2"], 50)

    def test_thousands_text_balance_and_existing_supply_are_added_correctly(self):
        self._build([100, 0, 1350, "-1,250"])
        _, cells = self._verify(1250, 0, {"I2", "K2"})
        self.assertEqual(cells["I2"], 1250)
        self._build([100, "1,250", 1450, -100])
        _, cells = self._verify(100, 0, {"I2", "K2"})
        self.assertEqual(cells["I2"], 1350)

    def test_supply_formula_preserves_its_expression_when_adding_quantity(self):
        self._build([100, "=20+5", 150, "=H2+I2-J2"])
        _, cells = self._verify(25, 0, {"I2"})
        self.assertEqual(cells["I2"], "=(20+5)+25")

    def test_legacy_formula_balance_preserves_references_without_touching_usage(self):
        self._build([100, 0, 125, "=H2+I2-J2"], legacy=True)
        result, cells = self._verify(25, 0, {"K2"})
        self.assertIsNone(result["supplement_col"])
        self.assertEqual(cells["K2"], "=(H2+I2-J2)+25")

    def test_uncalculable_cached_formula_refuses_write_and_preserves_original_bytes(self):
        self._build([100, 0, 125, "=SUM(H2:J2)"])
        self._cache_k2(-25)
        original = self.path.read_bytes()
        with self.assertRaises(ValueError):
            supplement_part_in_main(str(self.path), "PART-A", 25)
        self.assertEqual(self.path.read_bytes(), original)
        with self.assertRaisesRegex(ValueError, r"PART-A.*K2"):
            read_stock(str(self.path))

    def test_large_stock_cannot_hide_a_supplement_omitted_by_the_balance_formula(self):
        self._build([1_000_000_000, 0, 0, "=H2-J2"])
        original = self.path.read_bytes()
        with self.assertRaises(ValueError):
            supplement_part_in_main(str(self.path), "PART-A", 0.5)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(read_stock(str(self.path))["PART-A"], 1_000_000_000)

    def _build_deduction(self, *, formula_supply=False):
        self._build([100, "=0+0" if formula_supply else 0, 125,
                     "=H2+I2-J2" if formula_supply else -25,
                     5, "=K2-L2" if formula_supply else -30])
        wb = load_workbook(self.path)
        ws = wb.active
        ws["L1"] = "不良品扣帳"
        ws["M1"] = "結存"
        ws["N1"] = None
        wb.save(self.path)
        wb.close()

    def test_supplement_before_two_column_deduction_survives_future_recalculation(self):
        from app.services.main_file_recalc import recalc_batch_balances_for_cell

        for formula_supply in [False, True]:
            with self.subTest(formula_supply=formula_supply):
                self._build_deduction(formula_supply=formula_supply)
                result, cells = self._verify(30, 0, {"I2", "K2", "M2"})
                self.assertEqual(result["supplement_col"], 9)
                self.assertEqual(cells["L2"], 5)
                if formula_supply:
                    self.assertEqual(cells["I2"], "=(0+0)+30")
                    self.assertEqual(cells["K2"], "=H2+I2-J2")
                    self.assertEqual(cells["M2"], "=K2-L2")
                wb = load_workbook(self.path)
                ws = wb.active
                result = recalc_batch_balances_for_cell(ws, row=2, col=9)
                self.assertEqual(result["current_stock"], 0)
                if formula_supply:
                    self.assertEqual(ws["I2"].value, "=(0+0)+30")
                    self.assertEqual(ws["K2"].value, "=H2+I2-J2")
                    self.assertEqual(ws["M2"].value, "=K2-L2")
                ws["N1"], ws["O1"], ws["P1"] = "10-2", "用量", "結存"
                ws["N2"], ws["O2"], ws["P2"] = 0, 1, "=M2+N2-O2"
                later = recalc_batch_balances_for_cell(ws, row=2, col=9)
                self.assertEqual(later["current_stock"], -1)
                self.assertEqual(ws["P2"].value, "=M2+N2-O2")
                wb.save(self.path)
                wb.close()
                self.assertEqual(read_stock(str(self.path))["PART-A"], -1)

    def test_supplement_api_and_snapshot_use_the_verified_saved_stock(self):
        from app.models import SupplementPartRequest
        from app.routers import schedule

        self._build([100, 0, 125, "=H2+I2-J2"])
        with patch.object(schedule.db, "get_setting", return_value=str(self.path)), \
             patch.object(schedule, "BACKUP_DIR", Path(self.temp_dir.name) / "backups"), \
             patch.object(schedule.db, "get_manual_snapshot_moq", return_value={}), \
             patch.object(schedule.db, "save_snapshot") as save_snapshot, \
             patch.object(schedule.db, "set_setting"), \
             patch.object(schedule.db, "log_activity"):
            result = schedule.supplement_part(SupplementPartRequest(part_number="PART-A", supplement_qty=25))
        self.assertEqual(result["stock_before"], -25)
        self.assertEqual(result["stock_after"], 0)
        self.assertEqual(save_snapshot.call_args.args[0], read_stock(str(self.path)))
        self.assertEqual(save_snapshot.call_args.args[0]["PART-A"], 0)

    def test_supplement_api_reports_unreadable_formula_and_does_not_sync_snapshot(self):
        from fastapi import HTTPException
        from app.models import SupplementPartRequest
        from app.routers import schedule

        for values, cache, qty in [
            ([100, 0, 125, "=SUM(H2:J2)"], -25, 25),
            ([1_000_000_000, 0, 0, "=H2-J2"], None, 0.5),
        ]:
            with self.subTest(values=values):
                self._build(values)
                if cache is not None:
                    self._cache_k2(cache)
                original = self.path.read_bytes()
                with patch.object(schedule.db, "get_setting", return_value=str(self.path)), \
                     patch.object(schedule, "BACKUP_DIR", Path(self.temp_dir.name) / "backups"), \
                     patch.object(schedule, "refresh_snapshot_from_main") as sync_snapshot, \
                     patch.object(schedule.db, "log_activity") as log_activity:
                    with self.assertRaises(HTTPException) as raised:
                        schedule.supplement_part(SupplementPartRequest(part_number="PART-A", supplement_qty=qty))
                self.assertEqual(raised.exception.status_code, 422)
                self.assertIn("補料", raised.exception.detail)
                self.assertEqual(self.path.read_bytes(), original)
                sync_snapshot.assert_not_called()
                log_activity.assert_not_called()
