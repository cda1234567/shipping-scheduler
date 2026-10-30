from __future__ import annotations

import os
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from starlette.requests import Request

import app.database as db
from app.routers import dispatch as dispatch_router
from app.services import erp_purchase as ep
from app.services import low_stock, part_conversion as pc
from app.services import shortage_purchase as sp
from app.services.edit_auth import request_requires_edit_auth


def comp(part, qty=1.0, needed=10.0, desc=""):
    return {"part_number": part, "description": desc, "qty_per_board": qty, "scrap_factor": 0.0,
            "needed_qty": needed, "prev_qty_cs": 0.0, "is_dash": False, "is_customer_supplied": False}


class PureTests(unittest.TestCase):
    def test_transform_rename_when_no_b(self):
        out, merged = pc.transform_components([comp("a-1", 2, 20), comp("X", 1, 10)], "A-1", "B-1")
        self.assertFalse(merged)
        self.assertEqual([c["part_number"] for c in out], ["B-1", "X"])
        self.assertEqual(out[0]["qty_per_board"], 2)

    def test_transform_merge_sums_qty(self):
        src = [comp("A", 2, 20), comp("B", 3, 30)]
        out, merged = pc.transform_components(src, "A", "B")
        self.assertTrue(merged)
        self.assertEqual(len(out), 1)
        self.assertEqual((out[0]["part_number"], out[0]["qty_per_board"], out[0]["needed_qty"]), ("B", 5, 50))
        self.assertEqual(len(src), 2)  # 不改傳入資料

    def test_transform_untouched_when_no_a(self):
        out, merged = pc.transform_components([comp("X")], "A", "B")
        self.assertFalse(merged)
        self.assertEqual(out[0]["part_number"], "X")

    def test_resolve_chain_and_cycle(self):
        m = {"A": "B", "B": "C"}
        self.assertEqual(pc.resolve("a", m), "C")
        self.assertEqual(pc.resolve("Z", m), "Z")
        self.assertEqual(pc.resolve("A", {"A": "B", "B": "A"}), "B")
        self.assertEqual(pc.sources_by_target(m), {"C": ["A", "B"]})

    def test_effective_stock_with_negative_a(self):
        sources = {"B": ["A"]}
        total, own, merged = pc.effective_stock("b", {"B": 100.0, "A": -30.0}, sources)
        self.assertEqual((total, own), (70.0, 100.0))
        self.assertEqual(merged, [{"part": "A", "stock": -30.0}])
        self.assertEqual(pc.effective_stock("C", {}, sources), (0.0, 0.0, []))

    def test_signature_ignores_order(self):
        a = [comp("A", 1, 2), comp("B", 3, 4)]
        self.assertEqual(pc.components_signature(a), pc.components_signature(list(reversed(a))))
        self.assertNotEqual(pc.components_signature(a), pc.components_signature([comp("A", 1, 2)]))


class Base(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:", check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(db._CREATE_SQL)

        @contextmanager
        def temp_conn():
            try:
                yield self.conn
                self.conn.commit()
            except Exception:
                self.conn.rollback()
                raise

        p = patch.object(db, "get_conn", temp_conn)
        p.start()
        self.addCleanup(p.stop)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        for target in ("app.services.bom_editor._BOM_BACKUP_DIR", "app.services.bom_revision.BOM_HISTORY_DIR"):
            q = patch(target, self.dir / target.rsplit(".", 1)[1])
            q.start()
            self.addCleanup(q.stop)
        (self.dir / "_BOM_BACKUP_DIR").mkdir()
        self.st = {}
        q = patch.object(db, "get_st_inventory_stock", lambda: self.st)
        q.start()
        self.addCleanup(q.stop)

    def add_bom(self, bom_id, model, comps, filepath="", order_qty=10):
        db.save_bom_file({"id": bom_id, "filename": f"{bom_id}.xlsx", "filepath": filepath, "model": model,
                          "order_qty": order_qty, "uploaded_at": "2026-10-01", "components": comps})

    def write_xlsx(self, name, rows):
        """rows: [(qty, part, desc, needed)]，從第 5 列開始。"""
        path = self.dir / name
        wb = Workbook()
        ws = wb.active
        ws["H1"] = 1001
        ws["K1"] = 10
        ws["C2"] = "M1"
        ws["D2"] = "PCB"
        for i, (qty, part, desc, needed) in enumerate(rows, start=5):
            ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"F{i}"] = qty, part, desc, needed
        wb.save(path)
        wb.close()
        return str(path)

    def parts(self, bom_id):
        return sorted(c["part_number"] for c in db.get_bom_components(bom_id))


class PreviewTests(Base):
    def test_preview_counts_lists_and_flags_merge(self):
        self.add_bom("b1", "M1", [comp("A", 2, 20), comp("X")])
        self.add_bom("b2", "M2", [comp("A", 1, 10), comp("B", 4, 40)])
        self.add_bom("b3", "M3", [comp("X")])
        self.st = {"A": -5, "B": 50}
        r = pc.preview("a", "b")
        self.assertEqual((r["count"], r["merge_count"]), (2, 1))
        by_id = {x["bom_id"]: x for x in r["affected"]}
        self.assertFalse(by_id["b1"]["merge"])
        self.assertTrue(by_id["b2"]["merge"])
        self.assertEqual(by_id["b2"]["qty_per_board_after"], 5)
        self.assertEqual(by_id["b1"]["model"], "M1")
        self.assertEqual(r["effective_stock_b"], 45)
        self.assertEqual(r["db_only_count"], 2)
        # 預覽不改任何資料
        self.assertEqual(self.parts("b1"), ["A", "X"])
        self.assertEqual(db.get_active_part_conversions(), {})

    def test_preview_no_usage_warns(self):
        self.add_bom("b1", "M1", [comp("X")])
        r = pc.preview("A", "B")
        self.assertEqual(r["count"], 0)
        self.assertTrue(r["warnings"])

    def test_validation(self):
        for a, b in (("", "B"), ("A", ""), ("A", "a")):
            with self.assertRaises(pc.ConversionError):
                pc.preview(a, b)
        self.add_bom("b1", "M1", [comp("A")])
        pc.convert("A", "B")
        with self.assertRaises(pc.ConversionError):
            pc.preview("A", "C")  # A 已轉換
        with self.assertRaises(pc.ConversionError):
            pc.preview("C", "A")  # 目標是已停用的 A


class ConvertRestoreTests(Base):
    def test_db_only_round_trip(self):
        self.add_bom("b1", "M1", [comp("A", 2, 20, "電阻"), comp("X")])
        self.add_bom("b2", "M2", [comp("A", 1, 10), comp("B", 4, 40)])
        before1 = [pc._clean_component(c) for c in db.get_bom_components("b1")]
        r = pc.convert("A", "B", actor="tester")
        cid = r["conversion_id"]
        self.assertEqual(self.parts("b1"), ["B", "X"])
        self.assertEqual(self.parts("b2"), ["B"])
        merged = db.get_bom_components("b2")[0]
        self.assertEqual((merged["qty_per_board"], merged["needed_qty"]), (5, 50))
        self.assertEqual(db.get_active_part_conversions(), {"A": "B"})
        log = db.list_part_conversions()[0]
        self.assertEqual((log["part_a"], log["part_b"], log["actor"], log["status"]), ("A", "B", "tester", "active"))
        self.assertEqual({x["bom_id"] for x in log["affected"]}, {"b1", "b2"})
        self.assertTrue(any(l["action"] == "part_convert" for l in db.get_activity_logs()))

        res = pc.restore(cid, actor="tester2")
        self.assertEqual(res["restored_boms"], 2)
        self.assertEqual(db.get_active_part_conversions(), {})
        self.assertEqual(self.parts("b1"), ["A", "X"])
        self.assertEqual(self.parts("b2"), ["A", "B"])
        after1 = [pc._clean_component(c) for c in db.get_bom_components("b1")]
        self.assertEqual(sorted(after1, key=lambda c: c["part_number"]), sorted(before1, key=lambda c: c["part_number"]))
        self.assertEqual(db.get_part_conversion(cid)["status"], "restored")
        with self.assertRaises(pc.ConversionError):  # 不能重複還原
            pc.restore(cid)

    def test_dry_run_changes_nothing(self):
        self.add_bom("b1", "M1", [comp("A")])
        r = pc.convert("A", "B", dry_run=True)
        self.assertTrue(r["dry_run"])
        self.assertEqual(self.parts("b1"), ["A"])
        self.assertEqual(db.list_part_conversions(), [])

    def test_convert_without_usage_rejected(self):
        self.add_bom("b1", "M1", [comp("X")])
        with self.assertRaises(pc.ConversionError):
            pc.convert("A", "B")

    def test_transaction_rolls_back_on_failure(self):
        self.add_bom("b1", "M1", [comp("A")])
        real = db._replace_components

        def boom(conn, bom_id, components):
            real(conn, bom_id, components)
            raise RuntimeError("boom")

        with patch.object(db, "_replace_components", boom):
            with self.assertRaises(RuntimeError):
                pc.convert("A", "B")
        self.assertEqual(self.parts("b1"), ["A"])
        self.assertEqual(db.list_part_conversions(), [])
        self.assertEqual(db.get_active_part_conversions(), {})

    def test_restore_blocked_if_bom_edited_unless_force(self):
        self.add_bom("b1", "M1", [comp("A", 2, 20)])
        cid = pc.convert("A", "B")["conversion_id"]
        self.add_bom("b1", "M1", [comp("B", 9, 90)])  # 轉換後使用者又編輯
        with self.assertRaises(pc.ConversionError) as ctx:
            pc.restore(cid)
        self.assertEqual(ctx.exception.status_code, 409)
        self.assertEqual(self.parts("b1"), ["B"])
        pc.restore(cid, force=True)
        self.assertEqual(self.parts("b1"), ["A"])

    def test_restore_requires_newest_first(self):
        self.add_bom("b1", "M1", [comp("A")])
        c1 = pc.convert("A", "B")["conversion_id"]
        pc.convert("B", "C")
        with self.assertRaises(pc.ConversionError) as ctx:
            pc.restore(c1)
        self.assertEqual(ctx.exception.status_code, 409)
        self.assertEqual(self.parts("b1"), ["C"])

    def test_xlsx_rename_and_restore_file(self):
        path = self.write_xlsx("b1.xlsx", [(2, "A", "電阻", 20), (1, "X", "x", 10)])
        self.add_bom("b1", "M1", [comp("A", 2, 20, "電阻"), comp("X", 1, 10, "x")], filepath=path)
        original = Path(path).read_bytes()
        r = pc.convert("A", "B")
        self.assertEqual(r["affected"][0]["file_mode"], "xlsx")
        ws = load_workbook(path).active
        self.assertEqual([ws["C5"].value, ws["C6"].value], ["B", "X"])
        self.assertEqual(self.parts("b1"), ["B", "X"])
        pc.restore(r["conversion_id"])
        self.assertEqual(Path(path).read_bytes(), original)
        self.assertEqual(self.parts("b1"), ["A", "X"])

    def test_xlsx_merge_sums_and_clears_a_row(self):
        path = self.write_xlsx("b1.xlsx", [(2, "A", "電阻A", 20), (3, "B", "電阻B", 30)])
        self.add_bom("b1", "M1", [comp("A", 2, 20), comp("B", 3, 30)], filepath=path)
        r = pc.convert("A", "B")
        self.assertTrue(r["affected"][0]["merge"])
        ws = load_workbook(path).active
        self.assertEqual((ws["B6"].value, ws["F6"].value), (5, 50))
        self.assertIsNone(ws["C5"].value)
        comps = db.get_bom_components("b1")
        self.assertEqual([(c["part_number"], c["qty_per_board"]) for c in comps], [("B", 5.0)])

    def test_xlsx_failure_restores_file_and_db(self):
        # 用量是公式 → 無法自動合併 → 整筆失敗，檔案與資料庫都維持原樣
        path = self.write_xlsx("b1.xlsx", [("=1+1", "A", "a", 20), (3, "B", "b", 30)])
        self.add_bom("b1", "M1", [comp("A", 2, 20), comp("B", 3, 30)], filepath=path)
        original = Path(path).read_bytes()
        with self.assertRaises(pc.ConversionError):
            pc.convert("A", "B")
        self.assertEqual(Path(path).read_bytes(), original)
        self.assertEqual(self.parts("b1"), ["A", "B"])
        self.assertEqual(db.list_part_conversions(), [])


class ShortageIntegrationTests(Base):
    """A 轉 B 後：缺料採購清單不再出現 A；B 需求含 A、有效庫存 = B + A。"""

    def setUp(self):
        super().setUp()
        self.groups = [
            {"order_id": 1, "batch_code": "1-1", "model": "M", "po_number": "P1", "ship_date": "2026-12-28",
             "items": [{"part": "B-1", "desc": "x", "need_qty": 100}, {"part": "A-1", "desc": "x", "need_qty": 40}]},
        ]
        patches = [
            patch.object(dispatch_router, "build_dispatch_groups", lambda ids, dec: self.groups),
            patch.object(sp, "_erp_lookups", lambda parts, today: {
                "products": {}, "last_po": {}, "suppliers": {}, "on_order": {}, "warnings": [],
                "tab_stock": {}, "tab_stock_error": "", "tab_ware_ids": ["002"]}),
            patch.object(ep, "get_config", lambda: {}),
        ]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    def part_map(self):
        return {p["part_number"]: p for p in sp.build_shortage_list([1])["parts"]}

    def test_without_conversion_both_listed(self):
        self.st = {"B-1": 100, "A-1": -30}
        got = self.part_map()
        self.assertEqual(set(got), {"A-1", "B-1"})
        self.assertEqual(got["B-1"]["merged_from"], [])

    def test_a_hidden_and_b_uses_b_plus_a_stock(self):
        db.save_bom_file({"id": "b1", "filename": "b1.xlsx", "filepath": "", "model": "M", "order_qty": 1,
                          "uploaded_at": "x", "components": [comp("A-1")]})
        pc.convert("A-1", "B-1")
        self.st = {"B-1": 100, "A-1": 20}
        got = self.part_map()
        self.assertEqual(set(got), {"B-1"})
        b = got["B-1"]
        self.assertEqual(b["need"], 140)  # 100 + A 的 40
        self.assertEqual((b["st"], b["st_own"]), (120, 100))
        self.assertEqual(b["merged_from"], [{"part": "A-1", "stock": 20.0}])
        self.assertEqual(b["shortage"], 20)  # max(0, 140 - (100 + 20))
        self.assertEqual(b["suggested_qty"], 20)

    def test_negative_a_stock_reduces_effective_stock(self):
        db.apply_part_conversion("A-1", "B-1", [], {}, [], "t")
        self.st = {"B-1": 100, "A-1": -30}
        b = self.part_map()["B-1"]
        self.assertEqual(b["st"], 70)
        self.assertEqual(b["shortage"], 70)  # 140 - 70
        self.assertEqual(b["merged_from"][0]["stock"], -30)

    def test_effective_negative_clamps_like_existing_rule(self):
        db.apply_part_conversion("A-1", "B-1", [], {}, [], "t")
        self.st = {"B-1": 10, "A-1": -50}
        b = self.part_map()["B-1"]
        self.assertEqual(b["st"], -40)
        self.assertEqual(b["shortage"], 140)  # ST 為負時不加需求，沿用 summarize_requested_supply

    def test_restore_shows_a_again(self):
        db.save_bom_file({"id": "b1", "filename": "b1.xlsx", "filepath": "", "model": "M", "order_qty": 1,
                          "uploaded_at": "x", "components": [comp("A-1")]})
        cid = pc.convert("A-1", "B-1")["conversion_id"]
        pc.restore(cid)
        self.st = {"B-1": 100, "A-1": -30}
        self.assertEqual(set(self.part_map()), {"A-1", "B-1"})

    def test_allocation_context_uses_effective_stock(self):
        db.apply_part_conversion("A-1", "B-1", [], {}, [], "t")
        self.st = {"B-1": 100, "A-1": 20}
        ctx = sp._allocation_context([1])
        self.assertEqual(ctx["B-1"]["st"], 120)
        self.assertNotIn("A-1", ctx)


class RouterTests(Base):
    def setUp(self):
        super().setUp()
        os.environ.setdefault("PYTEST_CURRENT_TEST", "1")
        from main import app
        self.client = TestClient(app)

    def test_edit_auth_covers_posts_not_gets(self):
        def req(method, path):
            return Request({"type": "http", "method": method, "path": path, "headers": [], "query_string": b""})
        self.assertTrue(request_requires_edit_auth(req("POST", "/api/part-conversion/convert")))
        self.assertTrue(request_requires_edit_auth(req("POST", "/api/part-conversion/3/restore")))
        self.assertFalse(request_requires_edit_auth(req("GET", "/api/part-conversion/preview")))

    def test_api_preview_convert_restore(self):
        self.add_bom("b1", "M1", [comp("A", 2, 20)])
        self.add_bom("b2", "M2", [comp("A", 1, 10), comp("B", 4, 40)])
        r = self.client.get("/api/part-conversion/preview", params={"part_a": "a", "part_b": "b"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual((r.json()["count"], r.json()["merge_count"]), (2, 1))
        r = self.client.post("/api/part-conversion/convert", json={"part_a": "A", "part_b": "B", "dry_run": True})
        self.assertTrue(r.json()["dry_run"])
        self.assertEqual(self.client.get("/api/part-conversion/list").json()["active"], {})
        r = self.client.post("/api/part-conversion/convert", json={"part_a": "A", "part_b": "B"})
        self.assertEqual(r.status_code, 200, r.text)
        cid = r.json()["conversion_id"]
        lst = self.client.get("/api/part-conversion/list").json()
        self.assertEqual(lst["active"], {"A": "B"})
        self.assertEqual(lst["conversions"][0]["id"], cid)
        r = self.client.post(f"/api/part-conversion/{cid}/restore", json={})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.client.get("/api/part-conversion/list").json()["active"], {})
        self.assertEqual(self.client.post(f"/api/part-conversion/{cid}/restore", json={}).status_code, 409)

    def test_api_errors(self):
        self.assertEqual(self.client.get("/api/part-conversion/preview", params={"part_a": "A"}).status_code, 400)
        self.assertEqual(self.client.post("/api/part-conversion/999/restore", json={}).status_code, 404)


if __name__ == "__main__":
    unittest.main()
