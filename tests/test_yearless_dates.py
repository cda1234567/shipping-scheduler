from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

import app.database as db
from app.services import date_infer
from app.services.schedule_parser import _to_date_str, parse_schedule
from app.tools import fix_yearless_dates as fix

TODAY = date(2026, 10, 30)


class InferYearTests(unittest.TestCase):
    def test_window(self):
        self.assertEqual(date_infer.infer_year(1, 29, TODAY), date(2027, 1, 29))  # 今年 1/29 已過 > 60 天
        self.assertEqual(date_infer.infer_year(10, 1, TODAY), date(2026, 10, 1))  # 29 天前：仍是今年
        self.assertEqual(date_infer.infer_year(11, 15, TODAY), date(2026, 11, 15))
        self.assertEqual(date_infer.infer_year(8, 31, TODAY), date(2026, 8, 31))  # 60 天前整：不算「超過」
        self.assertEqual(date_infer.infer_year(8, 30, TODAY), date(2027, 8, 30))  # 61 天前
        self.assertEqual(date_infer.infer_year(12, 31, TODAY), date(2026, 12, 31))

    def test_feb_29(self):
        self.assertEqual(date_infer.infer_year(2, 29, date(2026, 12, 1)), date(2028, 2, 29))

    def test_format_detection(self):
        for fmt in ("m/d", "mm/dd", 'm"月"d"日"', "m-d", "d/m"):
            self.assertTrue(date_infer.format_is_yearless(fmt), fmt)
        for fmt in ("yyyy/m/d", "yy/m/d", "m/d/yyyy", "General", "", None, "0.00", "[$-404]e/m/d", "yyyy-mm-dd;@"):
            self.assertFalse(date_infer.format_is_yearless(fmt), str(fmt))
        # 引號內的文字不算格式記號
        self.assertFalse(date_infer.format_is_yearless('"year"'))

    def test_text(self):
        self.assertEqual(date_infer.parse_yearless_text("1/29", TODAY), date(2027, 1, 29))
        self.assertEqual(date_infer.parse_yearless_text("1-29", TODAY), date(2027, 1, 29))
        self.assertEqual(date_infer.parse_yearless_text("11月5日", TODAY), date(2026, 11, 5))
        self.assertIsNone(date_infer.parse_yearless_text("2027-01-29", TODAY))
        self.assertIsNone(date_infer.parse_yearless_text("13/40", TODAY))
        self.assertEqual(date_infer.normalize_date_text("1/29"), date_infer.infer_year(1, 29).isoformat())
        self.assertEqual(date_infer.normalize_date_text("2026-01-29"), "2026-01-29")


class ToDateStrTests(unittest.TestCase):
    def test_cell_values(self):
        dt = datetime(2026, 1, 29)
        self.assertEqual(_to_date_str(dt, "m/d", TODAY), "2027-01-29")
        self.assertEqual(_to_date_str(dt, "yyyy/m/d", TODAY), "2026-01-29")  # 明確年份不動
        self.assertEqual(_to_date_str(dt, "General", TODAY), "2026-01-29")
        self.assertEqual(_to_date_str(dt, "", TODAY), "2026-01-29")
        self.assertEqual(_to_date_str(datetime(2026, 11, 5), "m/d", TODAY), "2026-11-05")

    def test_strings(self):
        self.assertEqual(_to_date_str("1/29", "", TODAY), "2027-01-29")
        self.assertEqual(_to_date_str("1-29", "", TODAY), "2027-01-29")
        self.assertEqual(_to_date_str("2026/01/29", "", TODAY), "2026-01-29")
        self.assertEqual(_to_date_str("2026-01-29", "m/d", TODAY), "2026-01-29")
        self.assertIsNone(_to_date_str("abc", "", TODAY))
        self.assertIsNone(_to_date_str(None))


class ParseScheduleTests(unittest.TestCase):
    def test_parse_reads_number_format(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PCB"
        for i, (fmt, val) in enumerate([("m/d", datetime(2026, 1, 29)), ("yyyy/m/d", datetime(2026, 1, 29)),
                                        ("m/d", "1/29"), ("m/d", datetime(2026, 11, 20))], start=4):
            ws.cell(i, 1, 4500000000 + i)
            ws.cell(i, 2, "T1U")
            ws.cell(i, 3, "PCB1")
            ws.cell(i, 4, 100)
            ws.cell(i, 6, 10)
            c = ws.cell(i, 7, val)
            c.number_format = fmt
        with tempfile.TemporaryDirectory() as d:
            path = str(Path(d) / "s.xlsx")
            wb.save(path)
            rows = parse_schedule(path, today=TODAY)
        self.assertEqual([r.ship_date for r in rows], ["2027-01-29", "2026-01-29", "2027-01-29", "2026-11-20"])


class RequestModelTests(unittest.TestCase):
    def test_yearless_text_in_requests_gets_year(self):
        from app.models import PullinConfirmRequest, UpdateDeliveryRequest

        expect = date_infer.infer_year(1, 29).isoformat()
        self.assertEqual(UpdateDeliveryRequest(delivery_date="1/29").delivery_date, expect)
        self.assertEqual(UpdateDeliveryRequest(delivery_date="2026-01-29").delivery_date, "2026-01-29")
        self.assertEqual(PullinConfirmRequest(order_ids=[1], new_ship_date="1-29").new_ship_date, expect)


class FixToolTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:", check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(db._CREATE_SQL)

        @contextmanager
        def temp_conn():
            try:
                yield self.conn
                self.conn.commit()
            except Exception:
                self.conn.rollback()
                raise

        p = patch.object(db, "get_conn", temp_conn)
        p.start()
        self.addCleanup(p.stop)

    def add(self, code, status, ship, delivery, created):
        self.conn.execute(
            "INSERT INTO orders(code, po_number, model, status, ship_date, delivery_date, created_at) VALUES(?,?,?,?,?,?,?)",
            (code, "4500059664", "TA7", status, ship, delivery, created))

    def test_criterion(self):
        self.add("11-1", "dispatched", "2026-01-29", "2026-01-29", "2026-10-21T10:00:00")  # 要改
        self.add("3-1", "dispatched", "2026-03-06", "2026-03-06", "2026-03-11T09:00:00")  # 真正的舊單
        self.add("11-9", "cancelled", "2026-01-29", "2026-01-29", "2026-10-21T10:00:00")  # 已取消
        self.add("11-10", "completed", "2026-01-29", "2026-01-29", "2026-10-21T10:00:00")  # 已完成
        self.add("11-11", "pending", "2027-01-29", "2027-01-29", "2026-10-21T10:00:00")  # 本來就對
        self.add("11-12", "merged", "2026-01-29", "2026-10-30", "2026-10-21T10:00:00")  # 兩個欄位各自判斷
        self.add("11-13", "pending", "", None, "2026-10-21T10:00:00")
        changes = {c["code"]: c for c in fix.plan_fixes()}
        self.assertEqual(set(changes), {"11-1", "11-12"})
        self.assertEqual(changes["11-1"]["fields"]["ship_date"], {"old": "2026-01-29", "new": "2027-01-29"})
        self.assertEqual(set(changes["11-12"]["fields"]), {"ship_date"})
        # dry-run 不寫入
        fix.run()
        self.assertEqual(self.conn.execute("SELECT ship_date FROM orders WHERE code='11-1'").fetchone()[0], "2026-01-29")

    def test_apply_updates_and_logs_and_is_idempotent(self):
        self.add("11-1", "dispatched", "2026-01-29", "2026-01-29", "2026-10-21T10:00:00")
        self.add("3-1", "dispatched", "2026-03-06", "2026-03-06", "2026-03-11T09:00:00")
        self.add("old", "pending", "2024-01-29", "2024-01-29", "2026-10-21T10:00:00")  # 差兩年：加到不超過 60 天
        fix.run(apply=True)
        got = {r["code"]: (r["ship_date"], r["delivery_date"]) for r in self.conn.execute("SELECT * FROM orders")}
        self.assertEqual(got["11-1"], ("2027-01-29", "2027-01-29"))
        self.assertEqual(got["3-1"], ("2026-03-06", "2026-03-06"))
        self.assertEqual(got["old"], ("2027-01-29", "2027-01-29"))
        self.assertTrue(any(l["action"] == "fix_yearless_dates" for l in db.get_activity_logs()))
        self.assertEqual(fix.plan_fixes(), [])  # 再跑一次沒有東西可改

    def test_cli_defaults_to_dry_run(self):
        self.add("11-1", "dispatched", "2026-01-29", "2026-01-29", "2026-10-21T10:00:00")
        self.assertEqual(fix.main([]), 0)
        self.assertEqual(self.conn.execute("SELECT ship_date FROM orders").fetchone()[0], "2026-01-29")
        self.assertEqual(fix.main(["--apply"]), 0)
        self.assertEqual(self.conn.execute("SELECT ship_date FROM orders").fetchone()[0], "2027-01-29")


if __name__ == "__main__":
    unittest.main()
